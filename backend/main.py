"""
SAP Assessment Automation — FastAPI backend.

Start with:
    uvicorn main:app --host 127.0.0.1 --port 8000 --reload
"""

from __future__ import annotations

import io
import json
import logging
import threading
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Optional

import openpyxl
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from sap_connector import sap
from progress import ProgressTracker
from config import INPUT_DIR, OUTPUT_DIR, CREDENTIALS_FILE
import scoring
import transactions.taana as taana
import transactions.db15 as db15
import transactions.db02 as db02
import transactions.se16n as se16n
import transactions.se11 as se11
import transactions.aobj as aobj
import transactions.sara as sara

logging.basicConfig(level=logging.INFO, format="%(levelname)s | %(name)s | %(message)s")
logger = logging.getLogger(__name__)


@asynccontextmanager
async def lifespan(app: FastAPI):
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    logger.info("Input folder:  %s", INPUT_DIR)
    logger.info("Output folder: %s", OUTPUT_DIR)
    yield


app = FastAPI(
    title="SAP Assessment Automation API",
    version="1.0.0",
    description="Backend for SAP archivability analysis via GUI scripting.",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Single-user local app: one shared tracker per long-running batch operation
# is enough, no job IDs needed. The frontend polls the matching /progress
# endpoint while the background thread below runs the real work.
_db15_batch_progress = ProgressTracker()
_scoring_progress = ProgressTracker()


# ---------------------------------------------------------------------------
# Request / response models
# ---------------------------------------------------------------------------

class ConnectRequest(BaseModel):
    system: str
    client: str
    username: str
    password: str
    language: str = "EN"


class TaanaRequest(BaseModel):
    table_name: Optional[str] = None
    max_rows: int = 500


class Db15Request(BaseModel):
    table_name: str


class Se16nRequest(BaseModel):
    table_name: str
    max_rows: int = 200
    where_clause: Optional[str] = None


class Se11Request(BaseModel):
    table_name: str


class AobjRequest(BaseModel):
    object_filter: Optional[str] = None


class SaraRequest(BaseModel):
    archiving_object: str


class Db15ExportRequest(BaseModel):
    rows: list[dict[str, str]]


class Db15ScoreRequest(BaseModel):
    rows: list[dict[str, str]]


class Db15ScoreExportRequest(BaseModel):
    rows: list[dict[str, str]]
    recommended: list[dict[str, str]]


class Db02TopTablesRequest(BaseModel):
    limit: int = 300


class Db02ExportRequest(BaseModel):
    rows: list[dict[str, str]]


class Db02DebugClickRequest(BaseModel):
    tree_id: str
    node_key: str
    action: str = "select"


class Db15BatchFromInputRequest(BaseModel):
    filename: str


class SaveTablesRequest(BaseModel):
    rows: list[dict[str, str]]


class SaveCredentialsRequest(BaseModel):
    system: str
    client: str
    username: str
    password: str
    language: str = "EN"


# ---------------------------------------------------------------------------
# SAP session endpoints
# ---------------------------------------------------------------------------

@app.get("/api/health")
def health():
    return {"status": "ok", "sap_connected": sap.is_connected}


@app.get("/api/sap/systems")
def list_systems():
    """Return the SAP systems configured in the local SAP Logon pad."""
    systems = sap.list_systems()
    return {"systems": systems}


@app.post("/api/sap/connect")
def connect(req: ConnectRequest):
    result = sap.connect(
        system=req.system,
        client=req.client,
        username=req.username,
        password=req.password,
        language=req.language,
    )
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result["message"])
    return result


@app.post("/api/sap/disconnect")
def disconnect():
    return sap.disconnect()


@app.get("/api/sap/status")
def status():
    return {"connected": sap.is_connected}


@app.get("/api/sap/info")
def sap_info():
    """Return current connection state including system and user — used by the
    frontend on page load to restore the session without re-entering credentials."""
    return {
        "connected": sap.is_connected,
        "system": sap.connected_system,
        "user": sap.connected_user,
    }


@app.post("/api/sap/credentials")
def save_credentials(req: SaveCredentialsRequest):
    """Persist connection details (including password) to a local JSON file so
    the login form auto-fills on the next session."""
    data = {
        "system": req.system,
        "client": req.client,
        "username": req.username,
        "password": req.password,
        "language": req.language,
    }
    CREDENTIALS_FILE.write_text(json.dumps(data, indent=2), encoding="utf-8")
    return {"saved": True}


@app.get("/api/sap/credentials")
def load_credentials():
    """Return previously saved connection details, or an empty object if none."""
    if not CREDENTIALS_FILE.exists():
        return {}
    try:
        return json.loads(CREDENTIALS_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


# ---------------------------------------------------------------------------
# File folder endpoints
# ---------------------------------------------------------------------------

@app.get("/api/files/input")
def list_input_files():
    """List Excel files in the input folder, newest-modified first."""
    files = sorted(
        [f for f in INPUT_DIR.glob("*.xlsx")] + [f for f in INPUT_DIR.glob("*.xls")],
        key=lambda f: f.stat().st_mtime,
        reverse=True,
    )
    return {"files": [f.name for f in files]}


@app.post("/api/files/input/save")
def save_tables_to_input(req: SaveTablesRequest):
    """Save the table list rows to input/list_of_tables.xlsx."""
    columns = ["Table Name", "Description"]
    if req.rows and "Volume (GB)" in req.rows[0]:
        columns.append("Volume (GB)")
    buf = _build_workbook(req.rows, columns=columns, sheet_title="Top Tables")
    dest = INPUT_DIR / "list_of_tables.xlsx"
    dest.write_bytes(buf.getvalue())
    return {"saved": True, "path": str(dest)}


@app.post("/api/files/output/save-archiving")
def save_archiving_to_output(req: Db15ExportRequest):
    """Save archiving-objects rows to output/archiving_objects_by_table.xlsx."""
    buf = _build_workbook(
        req.rows,
        columns=["Table Name", "Table Description", "Archiving Object", "Object Description"],
        sheet_title="DB15 Results",
    )
    dest = OUTPUT_DIR / "archiving_objects_by_table.xlsx"
    dest.write_bytes(buf.getvalue())
    return {"saved": True, "path": str(dest)}


@app.post("/api/files/output/save-scored")
def save_scored_to_output(req: Db15ScoreExportRequest):
    """Save scored archiving-objects to output/archiving_objects_scored.xlsx."""
    buf = _build_multi_sheet_workbook([
        (
            "All Scored Objects",
            ["Table Name", "Table Description", "Archiving Object", "Object Description", "Score"],
            req.rows,
        ),
        (
            "Recommended",
            ["Table Name", "Table Description", "Archiving Object", "Object Description", "Score", "Rationale"],
            req.recommended,
        ),
    ])
    dest = OUTPUT_DIR / "archiving_objects_scored.xlsx"
    dest.write_bytes(buf.getvalue())
    return {"saved": True, "path": str(dest)}


# ---------------------------------------------------------------------------
# Transaction endpoints
# ---------------------------------------------------------------------------

@app.post("/api/transactions/taana")
def run_taana(req: TaanaRequest):
    _require_connection()
    result = taana.run(table_name=req.table_name, max_rows=req.max_rows)
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result["message"])
    return result


@app.post("/api/transactions/db15")
def run_db15(req: Db15Request):
    _require_connection()
    result = db15.run(table_name=req.table_name)
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result["message"])
    return result


@app.post("/api/transactions/se16n")
def run_se16n(req: Se16nRequest):
    _require_connection()
    result = se16n.run(
        table_name=req.table_name,
        max_rows=req.max_rows,
        where_clause=req.where_clause,
    )
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result["message"])
    return result


@app.post("/api/transactions/se11")
def run_se11(req: Se11Request):
    _require_connection()
    result = se11.run(table_name=req.table_name)
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result["message"])
    return result


@app.post("/api/transactions/aobj")
def run_aobj(req: AobjRequest):
    _require_connection()
    result = aobj.run(object_filter=req.object_filter)
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result["message"])
    return result


@app.post("/api/transactions/sara")
def run_sara(req: SaraRequest):
    _require_connection()
    result = sara.run(archiving_object=req.archiving_object)
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result["message"])
    return result


# ---------------------------------------------------------------------------
# DB15 batch (Excel upload -> run DB15 per table -> Excel export)
# ---------------------------------------------------------------------------

@app.post("/api/transactions/db15/batch")
def run_db15_batch(file: UploadFile = File(...)):
    _require_connection()
    if _db15_batch_progress.is_running():
        raise HTTPException(status_code=409, detail="A batch lookup is already in progress.")

    contents = file.file.read()
    tables = _parse_table_list(contents)
    if not tables:
        raise HTTPException(
            status_code=400,
            detail="No table names found in the uploaded file. Expected a table "
            "name in column A (and optionally a description in column B), "
            "starting from row 2.",
        )

    _db15_batch_progress.start(len(tables))

    def job():
        try:
            result = db15.run_batch(tables, on_progress=_db15_batch_progress.update)
            if result["status"] == "error":
                _db15_batch_progress.fail(result["message"])
            else:
                _db15_batch_progress.finish(result)
        except Exception as exc:
            logger.exception("DB15 batch job failed")
            _db15_batch_progress.fail(str(exc))

    threading.Thread(target=job, daemon=True).start()
    return {"status": "started", "total": len(tables)}


@app.post("/api/transactions/db15/batch-from-input")
def run_db15_batch_from_input(req: Db15BatchFromInputRequest):
    """Start a DB15 batch job using a file that already exists in the input folder."""
    _require_connection()
    if _db15_batch_progress.is_running():
        raise HTTPException(status_code=409, detail="A batch lookup is already in progress.")

    safe_name = Path(req.filename).name
    file_path = INPUT_DIR / safe_name
    if not file_path.exists():
        raise HTTPException(status_code=404, detail=f"'{safe_name}' not found in the input folder.")

    tables = _parse_table_list(file_path.read_bytes())
    if not tables:
        raise HTTPException(
            status_code=400,
            detail="No table names found in the file. Expected a table name in column A "
            "(and optionally a description in column B), starting from row 2.",
        )

    _db15_batch_progress.start(len(tables))

    def job():
        try:
            result = db15.run_batch(tables, on_progress=_db15_batch_progress.update)
            if result["status"] == "error":
                _db15_batch_progress.fail(result["message"])
            else:
                _db15_batch_progress.finish(result)
        except Exception as exc:
            logger.exception("DB15 batch job failed")
            _db15_batch_progress.fail(str(exc))

    threading.Thread(target=job, daemon=True).start()
    return {"status": "started", "total": len(tables)}


@app.get("/api/transactions/db15/batch/progress")
def get_db15_batch_progress():
    return _db15_batch_progress.snapshot()


@app.get("/api/transactions/db15/debug-screen")
def debug_db15_screen():
    """Diagnostic: dump every element on the DB15 selection screen (id, type,
    name, text) to discover the real radio-button / field IDs for this SAP
    system. Not used by the UI — call directly (e.g. via browser or curl)
    while connected to SAP."""
    _require_connection()
    result = db15.debug_dump_screen()
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result.get("message", "debug dump failed"))
    return result


@app.get("/api/transactions/db15/debug-grid")
def debug_db15_grid(table_name: str):
    """Diagnostic: filter DB15 by *table_name* and report exactly what
    happens reading the results grid (found / row count / column order /
    first cell), including raw error text on failure at each step."""
    _require_connection()
    result = db15.debug_read_grid(table_name)
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result.get("message", "debug dump failed"))
    return result


@app.post("/api/transactions/db15/export")
def export_db15_batch(req: Db15ExportRequest):
    buffer = _build_workbook(
        req.rows,
        columns=["Table Name", "Table Description", "Archiving Object", "Object Description"],
        sheet_title="DB15 Results",
    )
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=archiving_objects_by_table.xlsx"},
    )


@app.post("/api/transactions/db15/score")
def score_db15_batch(req: Db15ScoreRequest):
    if _scoring_progress.is_running():
        raise HTTPException(status_code=409, detail="A scoring run is already in progress.")

    distinct_tables = {row.get("Table Name", "") for row in req.rows}
    _scoring_progress.start(len(distinct_tables))

    def job():
        try:
            result = scoring.score_archiving_objects(req.rows, on_progress=_scoring_progress.update)
            if result["status"] == "error":
                _scoring_progress.fail(result["message"])
            else:
                _scoring_progress.finish(result)
        except Exception as exc:
            logger.exception("Scoring job failed")
            _scoring_progress.fail(str(exc))

    threading.Thread(target=job, daemon=True).start()
    return {"status": "started", "total": len(distinct_tables)}


@app.get("/api/transactions/db15/score/progress")
def get_db15_score_progress():
    return _scoring_progress.snapshot()


@app.post("/api/transactions/db15/score-export")
def export_db15_scored(req: Db15ScoreExportRequest):
    buffer = _build_multi_sheet_workbook([
        (
            "All Scored Objects",
            ["Table Name", "Table Description", "Archiving Object", "Object Description", "Score"],
            req.rows,
        ),
        (
            "Recommended",
            ["Table Name", "Table Description", "Archiving Object", "Object Description", "Score", "Rationale"],
            req.recommended,
        ),
    ])
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=archiving_objects_scored.xlsx"},
    )


# ---------------------------------------------------------------------------
# DB02 (generate a starting table list via the SQL Editor's "top tables by
# size" query, for users who don't already have a list of tables to check)
# ---------------------------------------------------------------------------

@app.post("/api/transactions/db02/top-tables")
def get_db02_top_tables(req: Db02TopTablesRequest):
    _require_connection()
    result = db02.run_get_top_tables(req.limit)
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result["message"])
    return result


@app.post("/api/transactions/db02/export")
def export_db02_top_tables(req: Db02ExportRequest):
    columns = ["Table Name", "Description"]
    if req.rows and "Volume (GB)" in req.rows[0]:
        columns.append("Volume (GB)")
    buffer = _build_workbook(
        req.rows,
        columns=columns,
        sheet_title="Top Tables",
    )
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=db02_top_tables.xlsx"},
    )


@app.get("/api/transactions/db02/debug-screen")
def debug_db02_screen(container_id: str = "wnd[0]/usr"):
    """Diagnostic: navigate to DB02 and dump every element under
    *container_id* (id, type, name, text). Not used by the UI — call
    directly while connected to SAP to discover the real tree/SQL editor
    structure before wiring up db02.run_get_top_tables()."""
    _require_connection()
    result = db02.debug_dump_screen(container_id)
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result.get("message", "debug dump failed"))
    return result


@app.get("/api/transactions/db02/debug-tree")
def debug_db02_tree(tree_id: str):
    """Diagnostic: navigate to DB02, find the tree control at *tree_id*, and
    list every node's key + display text, to identify the real node keys for
    "Diagnostics" and "SQL Editor" without guessing."""
    _require_connection()
    result = db02.debug_dump_tree(tree_id)
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result.get("message", "tree dump failed"))
    return result


@app.get("/api/transactions/db02/debug-probe-run")
def debug_db02_probe_run(limit: int = 20):
    """Diagnostic: drill into SQL Editor, try setting the query text,
    pressing F8 to execute, switching to the Result tab, and dumping the
    screen afterward — reporting each step's outcome independently."""
    _require_connection()
    result = db02.debug_probe_run(limit)
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result.get("message", "probe run failed"))
    return result


@app.post("/api/transactions/db02/debug-click")
def debug_db02_click(req: Db02DebugClickRequest):
    """Diagnostic: navigate to DB02, perform *action* on *node_key* in the
    tree at *tree_id*, then dump wnd[0]/usr afterwards so the effect of the
    click is visible in the response."""
    _require_connection()
    result = db02.debug_click_tree_node(req.tree_id, req.node_key, req.action)
    if result["status"] == "error":
        raise HTTPException(status_code=500, detail=result.get("message", "tree click failed"))
    return result


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _require_connection():
    if not sap.is_connected:
        raise HTTPException(
            status_code=403,
            detail="Not connected to SAP. POST /api/sap/connect first.",
        )


def _parse_table_list(contents: bytes) -> list[dict]:
    """Read (Table Name, Description) pairs from an uploaded Excel file.
    Table name is column A, description is column B; row 1 is a header."""
    workbook = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    sheet = workbook.active

    tables = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        name = str(row[0]).strip()
        description = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        tables.append({"table_name": name, "description": description})
    return tables


def _build_workbook(rows: list[dict], columns: list[str], sheet_title: str) -> io.BytesIO:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_title

    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(col, "") for col in columns])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _build_multi_sheet_workbook(sheets: list[tuple[str, list[str], list[dict]]]) -> io.BytesIO:
    """Build a workbook with one sheet per (title, columns, rows) tuple.
    Numeric-looking cell values (e.g. Score) are written as real numbers so
    they sort/filter correctly in Excel, instead of as text."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    for title, columns, rows in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append(columns)
        for row in rows:
            sheet.append([_numeric_or_raw(row.get(col, "")) for col in columns])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _numeric_or_raw(value):
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str) and value.strip():
        try:
            return int(value)
        except ValueError:
            try:
                return float(value)
            except ValueError:
                pass
    return value


if __name__ == "__main__":
    import uvicorn
    from config import API_HOST, API_PORT

    uvicorn.run("main:app", host=API_HOST, port=API_PORT, reload=False)
