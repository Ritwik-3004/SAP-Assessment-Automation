"""Retention Requirements: reading a retention schedule (Word table, prose read by a simulated model, long and cut-off
documents), matching archiving objects to its entries, the prompt pre-selection for small models, failure handling, the
saved workbook and the endpoints. A simulated model and the bundled sample schedule; no network, no SAP, no AI quota.
Prints PASS/FAIL lines and a final "N/N passed" (run by tests/test_suites.py)."""

import io, os, pathlib, re, sys, tempfile, time
from pathlib import Path
from unittest.mock import MagicMock
BACKEND = pathlib.Path(__file__).resolve().parents[2]
os.chdir(BACKEND)
sys.path.insert(0, str(BACKEND))
import openpyxl
import retention as rt

tmp = Path(tempfile.mkdtemp())
rt.results_folder = lambda task: tmp
fails = []
total = [0]


def check(label, ok, extra=""):
    total[0] += 1
    print(("PASS " if ok else "FAIL ") + label + (f"  [{extra}]" if extra else ""))
    if not ok:
        fails.append(label)


# ---------------- a fake AI model ----------------
state = {"budget": 14000, "calls": [], "daily": False}
rt.llm.load_settings = lambda: MagicMock()
rt.llm.not_configured_message = lambda s=None: None
rt.llm.budget = lambda s=None: {"reference_chars": state["budget"]}
rt.llm.max_workers = lambda s=None: 2
rt.llm.max_output_tokens = lambda s=None: 4000
rt.object_descriptions.resolve = lambda names, known=None, **k: {n: (known or {}).get(n) or "(description not found)" for n in names}

RULES = {"FI_DOCUMNT": "general ledger", "SD_VBAK": "pricing", "WORKITEM": "access requests", "BC_SBAL": "systems management"}


def fake_structured(system, user, schema, max_tokens=0, settings=None):
    state["calls"].append((schema.__name__, user))
    if state["daily"] and schema.__name__ == "_Choice":
        raise rt.llm.LLMRateLimitError("The Groq daily limit for this model has been reached.", daily=True)
    if schema.__name__ == "_Terms":
        obj = re.search(r"Archiving object: (\S+)", user).group(1)
        return schema(terms=[RULES.get(obj, "nothing")] + ["accounting", "ledger"])
    if schema.__name__ == "_Entries":
        return schema(entries=[
            {"record_series": "Tax invoices", "code": "", "description": "Sales and purchase invoices", "trigger": "end of tax year", "retention": "10 years", "basis": "Tax Act s.147"},
            {"record_series": "Personnel files", "code": "P-1", "description": "", "trigger": "", "retention": "7 years after termination", "basis": ""},
            {"record_series": "Notes without a period", "code": "", "description": "", "trigger": "", "retention": "", "basis": ""},
        ])
    obj = re.search(r"Archiving object: (\S+)", user).group(1)
    if obj == "BAD_NUMBER":
        return schema(entry=99, confidence="high", reason="x", also=[1, 77])
    want = RULES.get(obj)
    lines = [l for l in user.split("RETENTION SCHEDULE ENTRIES:\n")[1].splitlines()]
    for l in lines:
        n, rest = l.split(". ", 1)
        if want and want in rest.lower():
            others = [int(x.split(". ", 1)[0]) for x in lines if x is not l][:2]
            return schema(entry=int(n), confidence="high", reason=f"{obj} holds these records.", also=others)
    return schema(entry=0, confidence="low", reason="Nothing in the schedule covers it.", also=[])


rt.llm.structured = fake_structured
DOCX = open(BACKEND / "resources" / "Records-Retention-Schedule-Sample 1.docx", "rb").read()


def run(objects, documents):
    job = rt.RetentionJob()
    job.start(objects, documents)
    end = time.time() + 30
    while time.time() < end:
        s = job.snapshot()
        if s["status"] in ("done", "error"):
            return job, s
        time.sleep(0.05)
    raise TimeoutError(job.snapshot())


OBJS = [{"object": "FI_DOCUMNT", "description": "Financial accounting documents"}, {"object": "SD_VBAK", "description": ""},
        {"object": "WORKITEM", "description": "Work items"}, {"object": "ZZ_NONE", "description": "Something unrelated"}]

# ---- 1. the docx table, everything fits the prompt
job, s = run(OBJS, [(DOCX, "Records-Retention-Schedule-Sample 1.docx")])
rec = {r["object"]: r for r in s["records"]}
check("status done", s["status"] == "done" and s["completed"] == 4, s["message"])
check("document read as a table, 62 entries", s["documents"] == [{"name": "Records-Retention-Schedule-Sample 1.docx", "entries": 62, "method": "table"}], str(s["documents"]))
check("FI_DOCUMNT -> General Ledger, 10 Years, copied from the document", rec["FI_DOCUMNT"]["state"] == "done" and rec["FI_DOCUMNT"]["code"] == "FIN-002" and rec["FI_DOCUMNT"]["retention"] == "10 Years" and rec["FI_DOCUMNT"]["years"] == 10.0, str(rec["FI_DOCUMNT"]["retention"]))
check("reference note names file, table row, code, series, function", "table 1, row" in rec["FI_DOCUMNT"]["reference"] and "FIN-002 General Ledger" in rec["FI_DOCUMNT"]["reference"] and "Finance & Accounting" in rec["FI_DOCUMNT"]["reference"] and "Records-Retention" in rec["FI_DOCUMNT"]["reference"], rec["FI_DOCUMNT"]["reference"])
check("trigger event carried over", rec["FI_DOCUMNT"]["trigger"].startswith("Closure"))
check("SD_VBAK (no description) matched too", rec["SD_VBAK"]["code"] == "SAL-001" and rec["SD_VBAK"]["retention"] == "6 Years")
check("an event-based period is typed", rec["WORKITEM"]["retention"] == "2 Years" and rec["WORKITEM"]["period_type"] == "Fixed")
check("object nothing covers -> no_match with a note, no period invented", rec["ZZ_NONE"]["state"] == "no_match" and rec["ZZ_NONE"]["retention"] == "" and "No entry" in rec["ZZ_NONE"]["note"])
check("alternatives carry period and reference", len(rec["FI_DOCUMNT"]["alternatives"]) == 2 and rec["FI_DOCUMNT"]["alternatives"][0]["retention"] and "table 1, row" in rec["FI_DOCUMNT"]["alternatives"][0]["reference"])
check("all entries were shown to the model (no pre-filter) and no terms call was made", not any(c[0] == "_Terms" for c in state["calls"]))
check("description came from the file / resolver", rec["FI_DOCUMNT"]["description"] == "Financial accounting documents" and rec["ZZ_NONE"]["description"] == "Something unrelated")

wb = openpyxl.load_workbook(s["output_path"])
check("workbook: results sheet + entries read", wb.sheetnames == ["Retention requirements", "Entries read"], str(wb.sheetnames))
rows = list(wb["Retention requirements"].iter_rows(values_only=True))
check("results header names the Reference column", "Reference (where the period was found)" in rows[0] and rows[0][0] == "Archiving Object", str(rows[0][:4]))
fi = [r for r in rows if r[0] == "FI_DOCUMNT"][0]
check("results row: object, period, years, series, code, reference", fi[2] == "10 Years" and fi[3] == 10.0 and fi[6] == "General Ledger" and fi[7] == "FIN-002" and "row" in fi[10], str(fi[:11]))
check("entries sheet lists all 62", wb["Entries read"].max_row == 63)
check("file was saved into the task's folder", Path(s["output_path"]).parent == tmp and Path(s["output_path"]).name == "Retention requirements.xlsx")

# ---- 2. small prompt budget: the entries are pre-filtered by keywords from the model's terms
state["budget"] = 2500; state["calls"].clear()
job, s = run([OBJS[0]], [(DOCX, "d.docx")])
rec = {r["object"]: r for r in s["records"]}
listing = [c for c in state["calls"] if c[0] == "_Choice"][0][1]
shown = listing.split("RETENTION SCHEDULE ENTRIES:\n")[1].splitlines()
check("pre-filter: a terms call was made and fewer than 62 entries shown", any(c[0] == "_Terms" for c in state["calls"]) and 12 <= len(shown) < 62, str(len(shown)))
check("pre-filter keeps the relevant entry (FIN-002 General Ledger)", rec["FI_DOCUMNT"]["code"] == "FIN-002", str(rec["FI_DOCUMNT"]["code"]))
state["budget"] = 14000

# ---- 3. a model answer that names an entry that doesn't exist is not trusted
job, s = run([{"object": "BAD_NUMBER", "description": ""}], [(DOCX, "d.docx")])
r = s["records"][0]
check("invalid entry number -> no match; invalid 'also' numbers dropped", r["state"] == "no_match" and r["alternatives"] == [] , str(r["state"]))

# ---- 4. a document that is not a table: read by the model, entries located by page
TXT = ("=== Page 1 ===\nIntroduction to our policy.\n=== Page 2 ===\nTax invoices must be kept for 10 years.\n"
       "Personnel files are kept 7 years after termination.\n").encode()
state["calls"].clear()
job, s = run([{"object": "FI_DOCUMNT", "description": "Financial accounting documents"}], [(TXT, "policy.txt")])
ents = job.entries()
check("prose document read by the model: 2 entries (blank retention dropped)", [e["name"] for e in ents] == ["Tax invoices", "Personnel files"] and s["documents"][0]["method"] == "AI", str([e["name"] for e in ents]))
check("AI entries located in the text (page 2)", all(e["location"] == "Page 2" for e in ents), str([e["location"] for e in ents]))
check("AI entries are flagged in the reference note", "read by the AI model" in rt.reference_note(ents[0]) and "Page 2" in rt.reference_note(ents[0]), rt.reference_note(ents[0]))
check("regulation basis kept", ents[0]["basis"] == "Tax Act s.147")

# ---- 4b. a long prose document is read in parts; a part whose answer is cut off is split; a huge one is reported
NL = chr(10)
FILLER = [f"Line {i} of general policy text that says nothing about retention." for i in range(900)]
LONG = NL.join(["=== Page 1 ==="] + FILLER + ["=== Page 2 ===", "Tax invoices must be kept for 10 years.", ""]).encode()
state["calls"].clear()
job, s = run([OBJS[0]], [(LONG, "long.txt")])
ent_calls = [c for c in state["calls"] if c[0] == "_Entries"]
check("a long document is read in several parts", len(ent_calls) >= 2 and s["status"] == "done", str(len(ent_calls)))
check("entries from every part are merged without duplicates", [e["name"] for e in job.entries()] == ["Tax invoices", "Personnel files"], str([e["name"] for e in job.entries()]))
check("a complete read reports nothing missing", s["completeness"]["trimmed"] == [] and s["completeness"]["unprocessed"] == [])

orig = fake_structured
cut = {"n": 0}
def cutting(system, user, schema, max_tokens=0, settings=None):
    if schema.__name__ == "_Entries" and len(user) > 6000:
        cut["n"] += 1
        raise rt.llm.LLMError("The model's answer was cut off.")
    return orig(system, user, schema, max_tokens, settings)
rt.llm.structured = cutting
job, s = run([OBJS[0]], [(LONG, "long.txt")])
check("a part whose answer is cut off is split and asked again", cut["n"] >= 1 and s["status"] == "done" and len(job.entries()) == 2, f"cuts={cut['n']} entries={len(job.entries())}")
rt.llm.structured = orig

HUGE = NL.join(["=== Page 1 ==="] + [f"Line {i} of general policy text that says nothing about retention." for i in range(6000)]).encode()
job, s = run([OBJS[0]], [(HUGE, "huge.txt")])
check("a document longer than the parts that are read is reported as only partly read", any("huge.txt" in t["what"] for t in s["completeness"]["trimmed"]), str(s["completeness"]["trimmed"])[:160])
wbk = openpyxl.load_workbook(s["output_path"]) if s.get("output_path") else None
check("the workbook has a Notes sheet saying so", wbk is not None and "Notes" in wbk.sheetnames)

def limited(system, user, schema, max_tokens=0, settings=None):
    if schema.__name__ == "_Entries":
        raise rt.llm.LLMRateLimitError("The Groq daily limit for this model has been reached.", daily=True)
    return orig(system, user, schema, max_tokens, settings)
rt.llm.structured = limited
job, s = run([OBJS[0]], [(TXT, "policy.txt")])
check("a rate limit while reading stops the run with the model's message", s["status"] == "error" and "daily limit" in s["message"], s["message"])
rt.llm.structured = orig

# the small-prompt pre-selection is announced once, as a note
state["budget"] = 2500
job, s = run(OBJS[:2], [(DOCX, "d.docx")])
check("pre-selection is reported as a single note", len([n for n in s["completeness"]["notes"] if "pre-selected" in n]) == 1, str(s["completeness"]["notes"])[:200])
state["budget"] = 14000
job, s = run(OBJS[:1], [(DOCX, "d.docx")])
check("no note when every entry fits", s["completeness"]["notes"] == [])

# ---- 5. two documents; one unreadable
job, s = run([OBJS[0]], [(b"junk", "old.doc"), (DOCX, "a.docx"), (b"x", "pic.png")])
check("unreadable documents are skipped with warnings, the rest used", s["status"] == "done" and len(s["documents"]) == 1 and any("old.doc" in w for w in s["warnings"]) and any("pic.png" in w for w in s["warnings"]), str(s["warnings"]))

# ---- 6. nothing usable in the documents
job, s = run([OBJS[0]], [(b"just words, no retention rules", "empty.txt")])
rt.llm.structured = lambda *a, **k: rt._Entries(entries=[])
job, s = run([OBJS[0]], [(b"just words, no retention rules", "empty.txt")])
check("no entries found -> a clear error", s["status"] == "error" and "No retention entries" in s["message"], s["message"])
rt.llm.structured = fake_structured

# ---- 7. daily limit stops the run with a message and keeps what is saved
state["daily"] = True
job, s = run(OBJS, [(DOCX, "d.docx")])
check("daily rate limit stops the run with the model's message", s["status"] == "error" and "daily limit" in s["message"], s["message"])
check("workbook still written (entries read sheet)", Path(s["output_path"]).exists())
state["daily"] = False

# ---- 8. model not configured
rt.llm.not_configured_message = lambda s=None: "ANTHROPIC_API_KEY is not set."
job, s = run(OBJS, [(DOCX, "d.docx")])
check("model not configured -> error before reading", s["status"] == "error" and "API_KEY" in s["message"])
rt.llm.not_configured_message = lambda s=None: None

# ---- 9. HTTP endpoints
import main
from fastapi.testclient import TestClient
c = TestClient(main.app)
wbo = openpyxl.Workbook(); w1 = wbo.active; w1.title = "Recommended"; w1.append(["Table Name", "Archiving Object"]); w1.append(["BKPF", "FI_DOCUMNT"])
w2 = wbo.create_sheet("Grouped by Object"); w2.append(["Archiving Object", "Object Description"]); w2.append(["SD_VBAK", "Sales documents"]); w2.append(["FI_DOCUMNT", "FI docs"])
bio = io.BytesIO(); wbo.save(bio)
r = c.post("/api/retention/start", files=[("files", ("sched.docx", DOCX)), ("file", ("objs.xlsx", bio.getvalue()))])
check("start endpoint accepts documents + an objects file", r.status_code == 200 and r.json()["objects"] == ["SD_VBAK", "FI_DOCUMNT"], r.text[:120])
time.sleep(1.5)
p = c.get("/api/retention/progress").json()
check("progress endpoint", p["status"] == "done" and [x["object"] for x in p["records"]] == ["SD_VBAK", "FI_DOCUMNT"] and p["entry_count"] == 62, str(p["status"]))
check("entries endpoint", len(c.get("/api/retention/entries").json()["entries"]) == 62)
check("no documents -> 422/400", c.post("/api/retention/start", data={"filename": "x"}).status_code in (400, 422))
r = c.post("/api/retention/start", files=[("files", ("sched.docx", DOCX))])
check("no objects file -> 400", r.status_code == 400, r.text[:100])
check("cancel endpoint", c.post("/api/retention/cancel").json() == {"ok": True})
print(f"\n{total[0] - len(fails)}/{total[0]} passed")
