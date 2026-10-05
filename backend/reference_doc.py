"""
Reference document analysis.

Parses an uploaded document (Excel, PDF, PowerPoint) that contains past project
analysis of archiving object recommendations, then compares those mappings
against the current scored/recommended results to surface matches and mismatches.
"""

import io
import json
import re
import logging
from typing import Optional

import openpyxl
from pypdf import PdfReader

import llm
import object_descriptions

logger = logging.getLogger(__name__)


def _norm(value) -> str:
    """Canonical form of a SAP name for comparison: no whitespace, uppercase.
    SAP table/object names never contain spaces, so a space in a document is formatting noise."""
    return re.sub(r"\s+", "", str(value or "")).upper()


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def _extract_excel_text(contents: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"=== Sheet: {sheet_name} ===")
        for row in ws.iter_rows(values_only=True):
            cells = [str(v).strip() if v is not None else "" for v in row]
            if any(cells):
                lines.append(" | ".join(cells))
    wb.close()
    return "\n".join(lines)


def _extract_pdf_text(contents: bytes) -> str:
    reader = PdfReader(io.BytesIO(contents))
    lines: list[str] = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            lines.append(f"=== Page {i + 1} ===")
            lines.append(text)
    return "\n".join(lines)


def _extract_pptx_text(contents: bytes) -> str:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError as exc:
        raise ValueError(
            "python-pptx is required for PowerPoint files. "
            "Run: pip install python-pptx"
        ) from exc
    prs = Presentation(io.BytesIO(contents))
    lines: list[str] = []
    for i, slide in enumerate(prs.slides):
        lines.append(f"=== Slide {i + 1} ===")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                lines.append(shape.text)
    return "\n".join(lines)


def _extract_text(contents: bytes, filename: str) -> str:
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xls")):
        return _extract_excel_text(contents)
    if lower.endswith(".pdf"):
        return _extract_pdf_text(contents)
    if lower.endswith((".pptx", ".ppt")):
        return _extract_pptx_text(contents)
    raise ValueError(f"Unsupported file type: {filename}")


# ---------------------------------------------------------------------------
# Claude-based mapping extraction
# ---------------------------------------------------------------------------

def _extract_mappings_via_llm(text: str, table_names: list[str]) -> dict[str, str]:
    """
    Ask the active AI model (Claude or Groq, see llm.py) to extract SAP table →
    archiving object pairs from *text*. Returns {TABLE_NAME: ARCHIVING_OBJECT}
    (both uppercased). Raises llm.LLMError if the model can't be reached.
    """
    settings = llm.load_settings()
    char_limit = llm.budget(settings)["reference_chars"]

    tables_hint = ", ".join(table_names[:60])  # keep the prompt reasonable

    prompt = f"""You are analysing a reference document from an SAP archiving assessment project.
The document contains past experience or recommendations mapping SAP TABLE NAMES to ARCHIVING OBJECTS.

Extract every explicit mapping of a table name to an archiving object you can find.
Return ONLY a JSON object like:
{{"BKPF": "FI_DOCUMNT", "VBAK": "SD_VBAK", "MKPF": "MM_MATBEL"}}

We are particularly interested in these tables (but extract any you find):
{tables_hint}

Rules:
- Use exact SAP uppercase naming (e.g. FI_DOCUMNT, not fi_documnt).
- Only include mappings explicitly stated in the document — never guess.
- If a table appears with multiple objects, use the most recently mentioned one.
- Return ONLY the JSON object, no extra text.

Document:
{text[:char_limit]}
"""

    raw = llm.text("", prompt, max_tokens=2000, settings=settings)

    # Strip optional markdown fences
    if raw.startswith("```"):
        parts = raw.split("```")
        raw = parts[1].lstrip("json").strip() if len(parts) > 1 else raw

    try:
        mappings: dict = json.loads(raw)
        return {_norm(k): _norm(v) for k, v in mappings.items() if _norm(v)}
    except Exception:
        logger.warning("Could not parse Claude mapping response: %.200s", raw)
        return {}


# ---------------------------------------------------------------------------
# Comparison
# ---------------------------------------------------------------------------

def analyze_reference_doc(
    contents: bytes,
    filename: str,
    recommended: list[dict],
    known_descriptions: Optional[dict] = None,
) -> dict:
    """
    Parse *contents* (the uploaded reference document), extract table→object
    mappings, then compare with *recommended* (the current scored output).

    *known_descriptions* ({OBJECT: description} from the current run) feeds
    object_descriptions.resolve(), which also tries the remembered list, SAP (if connected) and
    the AI model, so every object the document proposes comes back with a description.

    A table the app found no archiving object for, which the document maps to one, is a
    mismatch (so the document can fill the gap), not a match.

    Returns
    -------
    {
        "status": "ok",
        "filename": str,
        "ref_mappings": {TABLE: OBJECT, ...},
        "annotated_recommended": [...rows with "Comments" and "Ref Doc Object" added...],
        "matches": [...rows where ref doc agrees...],
        "mismatches": [...rows where ref doc differs, with "Ref Doc Object" column...],
        "not_in_ref": [...rows not mentioned in the reference document...],
        "object_descriptions": {OBJECT: description, ...},  # for the objects the document proposes
    }
    """
    # 1. Extract text
    try:
        text = _extract_text(contents, filename)
    except ValueError as exc:
        return {"status": "error", "message": str(exc)}
    except Exception as exc:
        return {"status": "error", "message": f"Could not read file: {exc}"}

    if not text.strip():
        return {"status": "error", "message": "Could not extract any text from the document."}

    # 2. Get table names from current recommendations
    table_names = [
        _norm(row.get("Table Name", ""))
        for row in recommended
        if row.get("Table Name")
    ]

    # 3. Extract mappings using the active AI model
    try:
        ref_mappings = _extract_mappings_via_llm(text, table_names)
    except llm.LLMError as exc:
        return {"status": "error", "message": str(exc)}

    # 4. Compare
    matches: list[dict] = []
    mismatches: list[dict] = []
    not_in_ref: list[dict] = []
    annotated: list[dict] = []

    for row in recommended:
        table = _norm(row.get("Table Name", ""))
        current_obj = _norm(row.get("Archiving Object", ""))
        annotated_row = dict(row)

        if table in ref_mappings:
            ref_obj = ref_mappings[table]
            if ref_obj == current_obj:
                annotated_row["Comments"] = "Matches reference document"
                annotated_row["Ref Doc Object"] = ""
                matches.append(annotated_row)
            else:
                annotated_row["Comments"] = f"Reference document suggests: {ref_obj}"
                annotated_row["Ref Doc Object"] = ref_obj
                mismatches.append(annotated_row)
        else:
            annotated_row["Comments"] = ""
            annotated_row["Ref Doc Object"] = ""
            not_in_ref.append(annotated_row)

        annotated.append(annotated_row)

    # Every object the document proposes gets a description, so an override never leaves a blank.
    proposed = sorted({r["Ref Doc Object"] for r in mismatches if r.get("Ref Doc Object")})
    try:
        descriptions = object_descriptions.resolve(
            proposed,
            known={**object_descriptions.known_from_rows(recommended), **(known_descriptions or {})},
        )
    except Exception as exc:
        logger.warning("Could not resolve object descriptions: %s", exc, exc_info=True)
        descriptions = {}

    return {
        "status": "ok",
        "filename": filename,
        "ref_mappings": ref_mappings,
        "annotated_recommended": annotated,
        "matches": matches,
        "mismatches": mismatches,
        "not_in_ref": not_in_ref,
        "object_descriptions": descriptions,
    }


# ---------------------------------------------------------------------------
# Header tables: reference document maps ARCHIVING OBJECT -> HEADER TABLE
# ---------------------------------------------------------------------------

def _parse_mapping_json(raw: str) -> dict[str, str]:
    """Parse the model's {"KEY": "VALUE"} answer (tolerating ``` fences); {} if unusable."""
    raw = raw.strip()
    if raw.startswith("```"):
        raw = raw.strip("`")
        raw = raw[4:].strip() if raw.lower().startswith("json") else raw.strip()
    try:
        data = json.loads(raw)
        return {_norm(k): _norm(v) for k, v in data.items() if _norm(v)}
    except Exception:
        logger.warning("Could not parse header-table mapping response: %.200s", raw)
        return {}


def _extract_header_mappings_via_llm(text: str, object_names: list[str]) -> dict[str, str]:
    """Ask the active AI model for archiving object → header table pairs in *text*.
    Returns {ARCHIVING_OBJECT: HEADER_TABLE} (both uppercased). Raises llm.LLMError."""
    settings = llm.load_settings()
    char_limit = llm.budget(settings)["reference_chars"]
    objects_hint = ", ".join(object_names[:60])

    prompt = f"""You are analysing a reference document from an SAP archiving assessment project.
The document contains past experience mapping SAP ARCHIVING OBJECTS to their HEADER TABLE
(the main/root table of the archiving object, e.g. EKKO for MM_EKKO, BKPF for FI_DOCUMNT).

Extract every explicit mapping of an archiving object to its header table you can find.
Return ONLY a JSON object like:
{{"MM_EKKO": "EKKO", "FI_DOCUMNT": "BKPF", "SD_VBAK": "VBAK"}}

We are particularly interested in these archiving objects (but extract any you find):
{objects_hint}

Rules:
- Use exact SAP uppercase naming.
- Only include mappings explicitly stated in the document - never guess.
- Give ONE header table per archiving object; if several are mentioned, use the one the
  document calls the header (or the most recently mentioned).
- Return ONLY the JSON object, no extra text.

Document:
{text[:char_limit]}
"""
    raw = llm.text("", prompt, max_tokens=2000, settings=settings)
    return _parse_mapping_json(raw)


def analyze_header_reference(contents: bytes, filename: str, rows: list[dict]) -> dict:
    """
    Compare the app's header-table *rows* ({"Archiving Object", "Header Table", ...}) with an
    uploaded SME reference document (Excel, PDF or PowerPoint).

    Same result shape as analyze_reference_doc(), keyed on Archiving Object / Header Table:
    each annotated row gets "Reference Check" and "Ref Doc Header Table" (the app's own
    "Comments" column is left alone). A blank current Header Table with a reference value
    counts as a mismatch, so the reference document can fill it in.

    {"status": "ok", "filename", "ref_mappings", "annotated_rows", "matches", "mismatches",
     "not_in_ref"}  -- or {"status": "error", "message"}.
    """
    try:
        text = _extract_text(contents, filename)
    except ValueError as exc:
        return {"status": "error", "message": str(exc)}
    except Exception as exc:
        return {"status": "error", "message": f"Could not read file: {exc}"}

    if not text.strip():
        return {"status": "error", "message": "Could not extract any text from the document."}

    object_names = [_norm(r.get("Archiving Object", "")) for r in rows if r.get("Archiving Object")]
    try:
        ref_mappings = _extract_header_mappings_via_llm(text, object_names)
    except llm.LLMError as exc:
        return {"status": "error", "message": str(exc)}

    matches: list[dict] = []
    mismatches: list[dict] = []
    not_in_ref: list[dict] = []
    annotated: list[dict] = []

    for row in rows:
        obj = _norm(row.get("Archiving Object", ""))
        current = _norm(row.get("Header Table", ""))
        out = dict(row)

        ref = ref_mappings.get(obj)
        if ref is None:
            out["Reference Check"] = ""
            out["Ref Doc Header Table"] = ""
            not_in_ref.append(out)
        elif ref == current:
            out["Reference Check"] = "Matches reference document"
            out["Ref Doc Header Table"] = ""
            matches.append(out)
        else:
            out["Reference Check"] = f"Reference document suggests: {ref}"
            out["Ref Doc Header Table"] = ref
            mismatches.append(out)
        annotated.append(out)

    return {
        "status": "ok",
        "filename": filename,
        "ref_mappings": ref_mappings,
        "annotated_rows": annotated,
        "matches": matches,
        "mismatches": mismatches,
        "not_in_ref": not_in_ref,
    }
