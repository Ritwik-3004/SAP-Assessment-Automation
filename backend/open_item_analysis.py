"""
Open Item Analysis task: for each recommended archiving object, run the object's open-item query (from
resources/Open Closed Item details.xlsx, sheet "Open Item Analysis") in DB02's SQL Editor and collect the
results in output/Open item analysis.xlsx.

  * The objects come from the recommended-objects workbook saved by Find Archiving Objects (or any file
    the user chooses) -- the same input as the Archiving Object Analysis task.
  * The sheet holds one query per object, in a grid of cells that each read "OBJECT:" followed by the SQL.
    An object without a query is skipped with a note.
  * The sheet's queries list one row per document (year, month, category), which would be millions of rows,
    so by default they are wrapped to COUNT the items per year / month / category. The user sees exactly
    the query that runs and can edit it; an edited query is run as typed.
  * Only read-only statements are ever sent to SAP (SELECT / WITH).

The job runs on a background thread and talks to the UI through snapshot(). Every SAP call goes through
the single SAP worker thread, so a query being edited never touches the screen while another one runs;
an edit is queued and runs right after the current query.
"""

import logging
import re
import threading
import time
from collections import deque
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from config import OUTPUT_DIR
import transactions.db02 as db02
from sap_connector import sap

logger = logging.getLogger(__name__)

OUTPUT_NAME = "Open item analysis.xlsx"
QUERY_SHEET = Path(__file__).parent / "resources" / "Open Closed Item details.xlsx"
MAX_QUERY_SECONDS = 30 * 60   # give up waiting for one query after this long
MAX_PREVIEW_ROWS = 500

_sheet_cache: dict = {"mtime": None, "data": None}

# the columns the sheet's queries produce, and how the counted result names them
_STD_COLUMNS = ("Y", "M", "CAT")
_COLUMN_TITLES = {"Y": "Year", "M": "Month", "CAT": "Category", "ITEMS": "Items"}

_FORBIDDEN = (
    "INSERT", "UPDATE", "DELETE", "DROP", "ALTER", "TRUNCATE", "CREATE", "GRANT", "REVOKE", "CALL", "EXEC",
    "EXECUTE", "MERGE", "UPSERT", "REPLACE", "RENAME", "LOCK", "UNLOAD", "IMPORT", "EXPORT", "COMMIT", "ROLLBACK",
)
_FORBIDDEN_RE = re.compile(r"\b(" + "|".join(_FORBIDDEN) + r")\b")


# ---------------------------------------------------------------------------
# SQL helpers
# ---------------------------------------------------------------------------

def _scan(sql: str) -> tuple[str, str]:
    """Walk *sql* once and return (flat, code):
      flat -- one line: comments removed, whitespace runs outside quotes collapsed to one space;
      code -- the same, with the contents of string literals / quoted identifiers blanked out, so keywords can
              be looked for without matching text inside quotes."""
    flat: list[str] = []
    code: list[str] = []
    i, n = 0, len(sql)
    pending_space = False

    def emit(ch: str, blank: bool = False) -> None:
        nonlocal pending_space
        if pending_space and flat:
            flat.append(" ")
            code.append(" ")
        pending_space = False
        flat.append(ch)
        code.append(" " if blank else ch)

    while i < n:
        ch = sql[i]
        nxt = sql[i + 1] if i + 1 < n else ""
        if ch == "-" and nxt == "-":                      # line comment
            while i < n and sql[i] not in "\r\n":
                i += 1
            pending_space = True
        elif ch == "/" and nxt == "*":                    # block comment
            end = sql.find("*/", i + 2)
            i = n if end < 0 else end + 2
            pending_space = True
        elif ch in ("'", '"'):                            # string literal / quoted identifier
            emit(ch)
            i += 1
            while i < n:
                if sql[i] == ch:
                    if i + 1 < n and sql[i + 1] == ch:    # doubled quote = escaped quote
                        emit(ch, blank=True)
                        emit(ch, blank=True)
                        i += 2
                        continue
                    break
                emit(sql[i], blank=True)
                i += 1
            if i < n:
                emit(ch)
                i += 1
        elif ch.isspace():
            pending_space = True
            i += 1
        else:
            emit(ch)
            i += 1
    return "".join(flat).strip(), "".join(code).strip()


def validate_query(sql: str) -> str:
    """Raise ValueError unless *sql* is one read-only statement (SELECT / WITH)."""
    flat, code = _scan(sql or "")
    code = code.rstrip("; ").strip()
    if not code:
        raise ValueError("The query is empty.")
    if ";" in code:
        raise ValueError("Only one statement can be run at a time (remove the ';' between statements).")
    if not re.match(r"^\(*\s*(SELECT|WITH)\b", code, re.IGNORECASE):
        raise ValueError("Only SELECT queries can be run here.")
    found = _FORBIDDEN_RE.search(code.upper())
    if found:
        raise ValueError(f"The query contains '{found.group(1)}'. Only read-only queries can be run here.")
    return flat.rstrip("; ").strip()


def prepare_for_sap(sql: str) -> str:
    """The one-line, comment-free, validated form of *sql* that is typed into the SQL Editor."""
    return validate_query(sql)


def count_query(sql: str) -> str:
    """Wrap a sheet query that lists Y / M / CAT per document so it counts the documents per year, month and
    category. Queries that don't have that shape, or already aggregate, are returned unchanged."""
    text = (sql or "").strip().rstrip(";").strip()
    _, code = _scan(text)
    upper = code.upper()
    aliases = {m.group(1) for m in re.finditer(r"\bAS\s+([A-Z_][A-Z0-9_]*)\b", upper)}
    if not set(_STD_COLUMNS) <= aliases or re.search(r"\bGROUP\s+BY\b|\bCOUNT\s*\(", upper):
        return text
    return (
        "SELECT Y, M, CAT, COUNT(*) AS ITEMS\nFROM (\n" + text + "\n) T\nGROUP BY Y, M, CAT\nORDER BY Y, M, CAT"
    )


# a table's client column is usually MANDT, but BALHDR uses MANDANT and SWWWIHEAD uses CLIENT
_CLIENT_LITERAL = re.compile(r"(\b(?:MANDT|MANDANT|CLIENT)\s*=\s*)'(\d{3})'", re.IGNORECASE)


def use_client(sql: str, client: str) -> tuple[str, str]:
    """Replace the client the sheet's query hard-codes (MANDT = '300') with *client*, the one the user is logged on
    to. Returns (sql, the sheet's client), or (sql, "") when nothing needed changing (or the client is unknown)."""
    if not client:
        return sql, ""
    found = {m.group(2) for m in _CLIENT_LITERAL.finditer(sql)} - {client}
    if not found:
        return sql, ""
    return _CLIENT_LITERAL.sub(lambda m: f"{m.group(1)}'{client}'", sql), sorted(found)[0]


def query_kind(rec: dict, query: str) -> str:
    """Where a record's query comes from: "sheet" (the standard one, with the logged-on client), "sheet_client"
    (the sheet's own query, with the sheet's client -- the user reverted the substitution) or "edited"."""
    if same_query(query, rec["default_query"]):
        return "sheet"
    if rec.get("sheet_query") and same_query(query, rec["sheet_query"]):
        return "sheet_client"
    return "edited"


def same_query(a: str, b: str) -> bool:
    try:
        return _scan(a)[0].rstrip("; ") == _scan(b)[0].rstrip("; ")
    except Exception:
        return a == b


# ---------------------------------------------------------------------------
# The query sheet
# ---------------------------------------------------------------------------

_CELL = re.compile(r"^\s*([A-Za-z0-9_/]+)\s*:\s*(.+)$", re.DOTALL)


def load_queries() -> dict:
    """The queries in the sheet: {"queries": {OBJECT: {"sql", "area"}}, "problem": str}.

    The sheet (the one whose name contains "open item") is a grid of cells, each "OBJECT:" then the SQL; the
    text in the top row of a cell's column is its area (Sales Documents, FINANCE, ...). Re-read when the file
    changes. A missing or unreadable file gives an empty result with a "problem"."""
    try:
        mtime = QUERY_SHEET.stat().st_mtime
    except OSError:
        return {"queries": {}, "problem": f"The query sheet was not found ({QUERY_SHEET.name} in backend/resources)."}
    if _sheet_cache["mtime"] == mtime and _sheet_cache["data"] is not None:
        return _sheet_cache["data"]

    data: dict = {"queries": {}, "problem": ""}
    try:
        wb = openpyxl.load_workbook(QUERY_SHEET, read_only=False, data_only=True)
        try:
            ws = next((w for w in wb.worksheets if "open item" in w.title.lower()), None)
            if ws is None:
                data["problem"] = "The query sheet has no 'Open Item Analysis' sheet."
            else:
                for row in ws.iter_rows():
                    for cell in row:
                        if not isinstance(cell.value, str):
                            continue
                        m = _CELL.match(cell.value)
                        if not m:
                            continue
                        obj, sql = m.group(1).upper(), m.group(2).strip()
                        if not re.match(r"^\(*\s*(SELECT|WITH)\b", sql, re.IGNORECASE):
                            continue
                        header = ws.cell(row=1, column=cell.column).value
                        data["queries"].setdefault(obj, {"sql": sql, "area": str(header or "").strip()})
                if not data["queries"]:
                    data["problem"] = "No queries were found in the 'Open Item Analysis' sheet."
        finally:
            wb.close()
    except Exception as exc:
        logger.warning("Could not read the query sheet: %s", exc)
        data = {"queries": {}, "problem": f"The query sheet could not be read ({exc})."}
    _sheet_cache.update(mtime=mtime, data=data)
    return data


# ---------------------------------------------------------------------------
# The job
# ---------------------------------------------------------------------------

class OpenItemAnalysisJob:
    def __init__(self):
        self._lock = threading.RLock()
        self._reset()

    def _reset(self) -> None:
        self._state = {
            "status": "idle",        # idle | running | done
            "total": 0,
            "completed": 0,          # objects with an outcome (done / skipped / failed)
            "message": None,
            "current": None,         # the query being run: {"object", "query", "started"(epoch seconds)}
            "client_note": None,     # {"sheet_client", "client", "objects"} when the sheet's client was replaced
            "output_path": None,
            "warnings": [],
        }
        self._records: dict[str, dict] = {}
        self._order: list[str] = []
        self._results: dict[str, dict] = {}     # object -> {"columns": [...], "rows": [[...]]}
        self._queue: deque = deque()            # entries {"object", "query"} still to run, in order
        self._worker_alive = False
        self._cancel = False

    # -- public API -----------------------------------------------------------

    def snapshot(self) -> dict:
        with self._lock:
            state = dict(self._state)
            state["records"] = [dict(self._records[o]) for o in self._order]
            state["warnings"] = list(self._state["warnings"])
        sheet = load_queries()
        state["sheet"] = {"objects": sorted(sheet["queries"]), "problem": sheet["problem"]}
        return state

    def is_active(self) -> bool:
        with self._lock:
            return self._state["status"] == "running"

    def start(self, objects: list[dict]) -> None:
        """*objects*: [{"object", "description"}] in order. Raises ValueError if the query sheet is unusable and
        RuntimeError if a run is already going."""
        sheet = load_queries()
        if sheet["problem"]:
            raise ValueError(sheet["problem"])
        client = self._logon_client()
        with self._lock:
            if self._state["status"] == "running":
                raise RuntimeError("An open item analysis is already running.")
            self._reset()
            self._state.update(status="running", total=len(objects), message="Starting…")
            changed: list[str] = []
            sheet_clients: set[str] = set()
            for o in objects:
                name = o["object"]
                entry = sheet["queries"].get(name)
                original = count_query(entry["sql"]) if entry else ""
                # the sheet's queries name one client (300); run them in the client the user is logged on to
                adjusted, sheet_client = use_client(entry["sql"], client) if entry else ("", "")
                default = count_query(adjusted) if entry else ""
                if sheet_client:
                    changed.append(name)
                    sheet_clients.add(sheet_client)
                self._order.append(name)
                self._records[name] = {
                    "object": name, "description": o.get("description", ""), "area": entry["area"] if entry else "",
                    "state": "queued", "rows": 0, "note": "Waiting to run.", "query": default,
                    "default_query": default, "edited": False, "queued_edit": False, "seconds": 0,
                    "query_kind": "sheet",
                    # the substitution, so the user can see it and undo it
                    "sheet_query": original if sheet_client else "", "sheet_client": sheet_client,
                    "client": client if sheet_client else "",
                }
                self._queue.append({"object": name, "query": default or None})
            if changed:
                self._state["client_note"] = {
                    "sheet_client": ", ".join(sorted(sheet_clients)), "client": client, "objects": changed,
                }
            self._worker_alive = True
        threading.Thread(target=self._work, daemon=True, name="open-item-analysis").start()

    def submit_query(self, obj: str, query: str) -> None:
        """Run *query* (typed or edited by the user) for *obj*. If the object is still waiting its turn, its
        query is replaced; otherwise the run is queued to go next, after the query now running. Raises
        ValueError for an invalid / non-read-only query, KeyError for an object that is not in this run."""
        obj = obj.strip().upper()
        validate_query(query)
        with self._lock:
            if obj not in self._records:
                raise KeyError(obj)
            rec = self._records[obj]
            for entry in self._queue:
                if entry["object"] == obj:
                    entry["query"] = query
                    rec["query"] = query
                    rec["query_kind"] = query_kind(rec, query)
                    rec["edited"] = rec["query_kind"] == "edited"
                    rec["note"] = "Waiting to run (with your query)."
                    return
            self._queue.appendleft({"object": obj, "query": query, "edit": True})
            rec["queued_edit"] = True
            if self._state["status"] != "running":
                self._cancel = False
                self._state.update(status="running", message=f"Running your query for {obj}…")
            start_worker = not self._worker_alive
            self._worker_alive = True
        if start_worker:
            threading.Thread(target=self._work, daemon=True, name="open-item-analysis").start()

    def result(self, obj: str) -> Optional[dict]:
        """The collected result of one object (first rows only) for the preview, or None."""
        with self._lock:
            res = self._results.get(obj.strip().upper())
            if res is None:
                return None
            return {
                "columns": [_COLUMN_TITLES.get(c, c) for c in res["columns"]],
                "rows": res["rows"][:MAX_PREVIEW_ROWS],
                "total_rows": len(res["rows"]),
            }

    def cancel(self) -> None:
        """Stop after the query now running (SAP cannot be interrupted mid-query). Finished objects stay in the
        workbook."""
        with self._lock:
            self._cancel = True

    # -- plumbing -------------------------------------------------------------

    @staticmethod
    def _logon_client() -> str:
        try:
            return str(sap.run(lambda: sap.get_session().Info.Client) or "").strip()
        except Exception:
            return ""

    def _set(self, **changes) -> None:
        with self._lock:
            self._state.update(changes)

    def _count_done(self) -> int:
        return sum(1 for r in self._records.values() if r["state"] in ("done", "skipped", "failed"))

    def _work(self) -> None:
        while True:
            with self._lock:
                if self._cancel or not self._queue:
                    if self._cancel:
                        for entry in self._queue:
                            rec = self._records[entry["object"]]
                            if rec["state"] == "queued":
                                rec["state"], rec["note"] = "skipped", "Cancelled before it ran."
                            rec["queued_edit"] = False
                        self._queue.clear()
                    self._worker_alive = False
                    self._state.update(
                        status="done", current=None, completed=self._count_done(),
                        message="Cancelled." if self._cancel else "Finished.",
                    )
                    try:
                        self._write_workbook()
                    except Exception:
                        logger.exception("Could not write the Open item analysis workbook")
                    return
                item = self._queue.popleft()
            try:
                self._run_one(item)
            except Exception as exc:   # never let one object stop the run
                logger.exception("Open item analysis failed for %s", item["object"])
                self._finish_failed(item, f"Unexpected error: {exc}")
            self._save()

    def _save(self) -> None:
        try:
            self._write_workbook()
        except Exception as exc:
            logger.exception("Could not write the Open item analysis workbook")
            with self._lock:
                self._state["warnings"].append(f"Could not save the Excel file: {exc}")

    def _finish_failed(self, item: dict, note: str) -> None:
        with self._lock:
            rec = self._records[item["object"]]
            rec["queued_edit"] = False
            if item["object"] in self._results:   # an edited query that fails keeps the earlier result
                rec["state"] = "done"
                rec["note"] = f"{note.rstrip('. ')}. The previous result was kept."
            else:
                rec["state"], rec["rows"], rec["note"] = "failed", 0, note
            self._state.update(current=None, completed=self._count_done())

    def _run_one(self, item: dict) -> None:
        obj, query = item["object"], item.get("query")
        with self._lock:
            rec = self._records[obj]
        if not query:
            with self._lock:
                rec.update(state="skipped", rows=0, queued_edit=False, note=(
                    "No query available for this archiving object in the Open Closed Item details sheet - skipped."
                ))
                self._state.update(completed=self._count_done())
            return

        try:
            flat = prepare_for_sap(query)
        except ValueError as exc:
            self._finish_failed(item, f"Query not run: {exc}")
            return

        started = time.time()
        with self._lock:
            kind = query_kind(rec, query)
            rec.update(state="running", queued_edit=False, query=query, query_kind=kind,
                       edited=kind == "edited", note="Query running in SAP…")
            self._state.update(current={"object": obj, "query": query, "started": started},
                               message=f"Running the query for {obj}…")
        result = db02.run_query(flat, wait_seconds=MAX_QUERY_SECONDS)
        seconds = int(time.time() - started)
        if result.get("status") != "ok":
            with self._lock:
                rec["seconds"] = seconds
            self._finish_failed(item, f"The query failed: {result.get('message', 'unknown error')}")
            return

        with self._lock:
            self._results[obj] = {"columns": result["columns"], "rows": result["rows"]}
            rec.update(
                state="done", rows=len(result["rows"]), seconds=seconds,
                note=(f"{len(result['rows'])} row{'s' if len(result['rows']) != 1 else ''} in {seconds}s."
                      + (" Ran your edited query." if rec["query_kind"] == "edited"
                         else f" Ran the sheet's query with client {rec['sheet_client']} (as in the sheet)."
                         if rec["query_kind"] == "sheet_client" else "")
                      + (f" The result reached the {len(result['rows']):,}-row limit and may be incomplete."
                         if result.get("truncated") else "")),
            )
            self._state.update(current=None, completed=self._count_done())

    # -- output ---------------------------------------------------------------

    @staticmethod
    def _cell(column: str, value: str):
        """Counts become numbers; everything else stays as SAP shows it."""
        if column.upper() in ("ITEMS", "COUNT", "CNT"):
            digits = "".join(ch for ch in str(value) if ch.isdigit())
            if digits:
                return int(digits)
        return value

    def _write_workbook(self) -> None:
        with self._lock:
            order = list(self._order)
            records = {o: dict(self._records[o]) for o in order}
            results = {o: {"columns": list(r["columns"]), "rows": list(r["rows"])} for o, r in self._results.items()}

        wb = openpyxl.Workbook()
        data = wb.active
        data.title = "Open item analysis"
        columns: list[str] = []                   # every result column, in order of first appearance
        for o in order:
            for c in results.get(o, {}).get("columns", []):
                if c not in columns:
                    columns.append(c)
        data.append(["Archiving Object"] + [_COLUMN_TITLES.get(c, c) for c in columns])
        for o in order:
            res = results.get(o)
            if not res:
                continue
            index = {c: i for i, c in enumerate(res["columns"])}
            for row in res["rows"]:
                data.append([o] + [
                    self._cell(c, row[index[c]]) if c in index and index[c] < len(row) else "" for c in columns
                ])
        self._style_header(data)
        data.freeze_panes = "B2"
        self._autosize(data)

        summary = wb.create_sheet("Summary")
        summary.append(["Archiving Object", "Description", "Area", "Result", "Rows", "Query", "Note", "SQL run"])
        for o in order:
            r = records[o]
            summary.append([
                o, r["description"], r["area"], r["state"].capitalize(), r["rows"] if r["state"] == "done" else "",
                self._query_source(r),
                r["note"], r["query"] if r["state"] in ("done", "failed") else "",
            ])
        self._style_header(summary)
        self._autosize(summary)
        summary.column_dimensions["H"].width = 80
        for row in summary.iter_rows(min_row=2):
            row[7].alignment = Alignment(wrap_text=False, vertical="top")

        path = self._save_workbook(wb)
        self._set(output_path=str(path))

    @staticmethod
    def _query_source(r: dict) -> str:
        if r["state"] in ("done", "failed"):
            if r["query_kind"] == "edited":
                return "Edited by you"
            if r["query_kind"] == "sheet_client":
                return f"From the sheet (client {r['sheet_client']}, as in the sheet)"
            if r["sheet_client"]:
                return f"From the sheet (client {r['client']} instead of {r['sheet_client']})"
            return "From the sheet"
        if r["state"] == "skipped" and not r["default_query"]:
            return "No query in the sheet"
        return ""

    @staticmethod
    def _style_header(ws) -> None:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center")

    @staticmethod
    def _autosize(ws) -> None:
        widths: dict[int, int] = {}
        for row in ws.iter_rows(values_only=True):
            for c_idx, value in enumerate(row, start=1):
                first_line = str(value).splitlines()[0] if value is not None and str(value) else ""
                widths[c_idx] = max(widths.get(c_idx, 8), min(len(first_line) + 2, 60))
        for c_idx, width in widths.items():
            ws.column_dimensions[get_column_letter(c_idx)].width = width

    @staticmethod
    def _save_workbook(wb) -> Path:
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        stem, suffix = Path(OUTPUT_NAME).stem, Path(OUTPUT_NAME).suffix
        for path in [OUTPUT_DIR / OUTPUT_NAME] + [OUTPUT_DIR / f"{stem} ({n}){suffix}" for n in range(2, 10)]:
            try:
                wb.save(path)
                return path
            except PermissionError:   # open in Excel
                continue
        raise PermissionError("Could not write the Open item analysis workbook (is it open in Excel?).")


job = OpenItemAnalysisJob()
