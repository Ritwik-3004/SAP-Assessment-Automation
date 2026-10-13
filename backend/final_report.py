"""
Final report generation: combine the important result sheets of the other tasks into two workbooks.

  * "Final report - Overview.xlsx" -- sheet "Overview": one block per archiving object (largest cumulative
    size first): archiving object, description, header table, its tables, archiving conditions, the objects to
    archive before it and the four size columns. Whatever is common to a group (object, description, header table, conditions, cumulative
    sizes) is merged into one cell across the group's table rows, like the "Grouped by Object" sheet.
    Sheet "Housekeeping Programs": the groups of that sheet that have a housekeeping program instead of an
    archiving object -- program, its tables and the four size columns, merged the same way.
  * "Final report - Object details.xlsx" -- one sheet per archiving object, holding that object's Table
    analysis (the sheet of its header table) and, next to it, its Open item analysis if there is one.

Both are written to results/Final report generation/ by the caller (main.py passes its save function, which
makes a numbered copy when a file is open in Excel). Nothing here talks to SAP: it only reads the workbooks
the earlier tasks saved.

Where the sources come from:
  objects, descriptions, housekeeping programs, tables, sizes
                                        results/Find Archiving Objects for Tables/archiving_objects_with_reference.xlsx
                                        (else archiving_objects_scored.xlsx), sheet "Grouped by Object"
  header table                          results/Find Header Tables for Archiving Objects/header_tables_with_reference.xlsx
                                        (else header_tables.xlsx)
  archiving conditions, archive before  results/Archiving Object Analysis/Archiving object analysis.xlsx
                                        (conditions: each object's sheet; archive before: the Summary sheet)
  table analysis                        results/Table Analysis/Table analysis.xlsx
  open item analysis                    results/Open Item Analysis/Open item analysis.xlsx
"""

from __future__ import annotations

import io
import math
import re
from pathlib import Path
from typing import Callable, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import results_folder

OVERVIEW_NAME = "Final report - Overview.xlsx"
DETAILS_NAME = "Final report - Object details.xlsx"


class ReportError(Exception):
    """Something the user can fix (a missing source file); the message says what."""


# ---------------------------------------------------------------------------
# Finding the source files
# ---------------------------------------------------------------------------

def _newest(folder: Path, stem: str) -> Optional[Path]:
    """The most recently written "<stem>*.xlsx" (a numbered copy such as "<stem> (2).xlsx" counts: it is what
    was written while the original was open in Excel)."""
    files = [p for p in folder.glob(f"{stem}*.xlsx") if not p.name.startswith("~$")]
    return max(files, key=lambda p: p.stat().st_mtime) if files else None


def _pair(folder: Path, preferred: str, fallback: str) -> tuple[Optional[Path], str]:
    """(file to use, warning): the "with reference" file if it exists, else the plain one."""
    first, second = folder / preferred, folder / fallback
    if first.exists():
        warning = ""
        if second.exists() and second.stat().st_mtime > first.stat().st_mtime:
            warning = (f"{preferred} is older than {fallback}: it may come from an earlier run. Run the "
                       "reference-document review again, or delete it to use the plain file.")
        return first, warning
    return (second if second.exists() else None), ""


def resolve_sources() -> dict:
    """{key: {"label", "path": Path|None, "fallback": bool, "warning": str, "required": bool}}."""
    objects, objects_warning = _pair(results_folder("BATCH"), "archiving_objects_with_reference.xlsx", "archiving_objects_scored.xlsx")
    headers, headers_warning = _pair(results_folder("HEADER_TABLES"), "header_tables_with_reference.xlsx", "header_tables.xlsx")
    return {
        "objects": {"label": "Archiving objects, descriptions, tables and sizes", "path": objects, "warning": objects_warning,
                    "fallback": bool(objects and "with_reference" not in objects.name), "required": True},
        "headers": {"label": "Header tables", "path": headers, "warning": headers_warning,
                    "fallback": bool(headers and "with_reference" not in headers.name), "required": False},
        "conditions": {"label": "Archiving conditions and archive-before objects", "required": False, "fallback": False, "warning": "",
                       "path": _newest(results_folder("OBJECT_ANALYSIS"), "Archiving object analysis")},
        "table_analysis": {"label": "Table analysis", "required": False, "fallback": False, "warning": "",
                           "path": _newest(results_folder("TABLE_ANALYSIS"), "Table analysis")},
        "open_items": {"label": "Open item analysis", "required": False, "fallback": False, "warning": "",
                       "path": _newest(results_folder("OPEN_ITEM_ANALYSIS"), "Open item analysis")},
    }


def describe_sources() -> list[dict]:
    """The sources for the panel: what will be used, and what is missing."""
    out = []
    for src in resolve_sources().values():
        path: Optional[Path] = src["path"]
        if path is None:
            status = "missing"
            note = "Required." if src["required"] else "Optional: that part of the report is left out."
        elif src["fallback"]:
            status, note = "fallback", "The reference-document file was not found, so the plain file is used."
        else:
            status, note = "found", ""
        if src["warning"]:
            note = (note + " " + src["warning"]).strip()
        out.append({"label": src["label"], "file": path.name if path else "", "folder": str(path.parent) if path else "",
                    "status": status, "note": note})
    return out


# ---------------------------------------------------------------------------
# Reading the source workbooks
# ---------------------------------------------------------------------------

def _load(path: Path):
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except PermissionError:
        raise ReportError(f"Could not read {path.name} — is it open in Excel? Close it and try again.")
    except Exception as exc:
        raise ReportError(f"Could not read {path.name}: {exc}")


def _fill_merged(ws) -> None:
    """Give every cell of a merged range the merged value, so a group's cells can be read row by row."""
    for rng in list(ws.merged_cells.ranges):
        value = ws.cell(rng.min_row, rng.min_col).value
        ws.unmerge_cells(str(rng))
        for row in ws.iter_rows(min_row=rng.min_row, max_row=rng.max_row, min_col=rng.min_col, max_col=rng.max_col):
            for cell in row:
                cell.value = value


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _grouped_rows(path: Path) -> list[tuple]:
    """The data rows (columns A-I) of the "Grouped by Object" sheet, merged cells filled in. Columns by position,
    as in the sheet: A object, B description, C housekeeping program, D table, F/G volume, H/I cumulative."""
    wb = _load(path)
    try:
        titled = [ws for ws in wb.worksheets if "grouped" in ws.title.lower()]
        ws = titled[-1] if titled else wb.worksheets[-1]
        if _text(ws.cell(1, 1).value).lower() != "archiving object":
            raise ReportError(f"{path.name} has no \"Grouped by Object\" sheet (its last sheet should start with an "
                              "\"Archiving Object\" column). Run Find Archiving Objects for Tables again.")
        _fill_merged(ws)
        return list(ws.iter_rows(min_row=2, max_col=9, values_only=True))
    finally:
        wb.close()


def read_groups(path: Path) -> list[dict]:
    """The archiving objects of the "Grouped by Object" sheet, in its order (largest first):
    [{"object", "description", "cum_gb", "cum_mb", "tables": [{"table", "gb", "mb"}]}].
    Groups without an archiving object (housekeeping programs only, or nothing found) are left out."""
    groups: dict[str, dict] = {}
    for row in _grouped_rows(path):
        obj = _text(row[0]).upper()
        if not obj or obj.startswith("("):
            continue
        g = groups.setdefault(obj, {"object": obj, "description": "", "cum_gb": None, "cum_mb": None, "tables": []})
        g["description"] = g["description"] or _text(row[1])
        if g["cum_gb"] is None:
            g["cum_gb"], g["cum_mb"] = row[7], row[8]
        if _text(row[3]):
            g["tables"].append({"table": _text(row[3]), "gb": row[5], "mb": row[6]})
    return list(groups.values())


def read_housekeeping_groups(path: Path) -> list[dict]:
    """The housekeeping programs of the same sheet, in its order: [{"program", "cum_gb", "cum_mb", "tables"}].
    These are the groups with no archiving object (a table has an archiving object or a housekeeping program).
    A program shown on an archiving object's rows is not one of them."""
    groups: dict[str, dict] = {}
    for row in _grouped_rows(path):
        program = _text(row[2])
        if _text(row[0]) or not program:
            continue
        g = groups.setdefault(program, {"program": program, "cum_gb": None, "cum_mb": None, "tables": []})
        if g["cum_gb"] is None:
            g["cum_gb"], g["cum_mb"] = row[7], row[8]
        if _text(row[3]):
            g["tables"].append({"table": _text(row[3]), "gb": row[5], "mb": row[6]})
    return list(groups.values())


def read_header_tables(path: Path) -> dict[str, str]:
    """{ARCHIVING OBJECT: header table} from the final list: the "(with Reference)" sheet if there is one,
    else "Header Tables" (column A object, column B header table)."""
    wb = _load(path)
    try:
        sheets = wb.worksheets
        ws = next((s for s in sheets if "with reference" in s.title.lower()), None) \
            or next((s for s in sheets if s.title.lower() == "header tables"), None) or sheets[-1]
        out: dict[str, str] = {}
        for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            obj, table = _text(row[0]).upper(), _text(row[1]).upper()
            if obj and table and obj not in out:
                out[obj] = table
        return out
    finally:
        wb.close()


def read_conditions(path: Path) -> dict[str, list[str]]:
    """{ARCHIVING OBJECT: [condition, ...]} from the object sheets of the analysis workbook (the Condition
    column of the table under the "ARCHIVING CONDITIONS" band: No. / Condition / Source / Detail / Reference check)."""
    wb = _load(path)
    try:
        out: dict[str, list[str]] = {}
        for ws in wb.worksheets:
            title = _text(ws.cell(1, 1).value)
            if not title.lower().startswith("archiving object:"):
                continue
            obj = title.split(":", 1)[1].strip().upper()
            rows = list(ws.iter_rows(values_only=True))
            start = next((i for i, r in enumerate(rows) if _text(r[0]).upper() == "ARCHIVING CONDITIONS"), None)
            conditions: list[str] = []
            if start is not None and start + 1 < len(rows) and _text(rows[start + 1][0]) == "No.":
                for r in rows[start + 2:]:
                    if not isinstance(r[0], (int, float)):
                        break
                    r = list(r) + [None] * 5
                    if _text(r[1]):
                        conditions.append(_text(r[1]))
            out[obj] = conditions
        return out
    finally:
        wb.close()


def read_archive_before(path: Path) -> dict[str, str]:
    """{ARCHIVING OBJECT: the objects to archive before it, as one cell of text} from the analysis workbook's
    Summary sheet (column "ARCHIVE BEFORE (in order)": the objects in archiving order, comma separated, or
    "None", or a note that the network was not available). A list becomes numbered lines, step 1 first."""
    wb = _load(path)
    try:
        ws = wb["Summary"] if "Summary" in wb.sheetnames else None
        if ws is None:
            return {}
        header = [_text(h).lower() for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
        obj_col = next((i for i, h in enumerate(header) if h == "archiving object"), None)
        ab_col = next((i for i, h in enumerate(header) if h.startswith("archive before")), None)
        if obj_col is None or ab_col is None:
            return {}
        out: dict[str, str] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            obj = _text(row[obj_col]).upper() if obj_col < len(row) else ""
            raw = _text(row[ab_col]) if ab_col < len(row) else ""
            if not obj or not raw:
                continue
            if raw.lower() == "none" or raw.startswith("("):
                out[obj] = raw
            else:
                items = [i.strip() for i in raw.split(",") if i.strip()]
                out[obj] = "\n".join(f"{n}. {i}" for n, i in enumerate(items, start=1))
        return out
    finally:
        wb.close()


def _grid(ws, header_row: int, first_col: int = 1) -> tuple[list[str], list[list]]:
    """Header and data rows of a sheet's table, up to the first fully blank row."""
    rows = list(ws.iter_rows(min_row=header_row, min_col=first_col, values_only=True))
    if not rows:
        return [], []
    columns = [_text(c) for c in rows[0]]
    while columns and not columns[-1]:
        columns.pop()
    data = []
    for r in rows[1:]:
        r = list(r)[: len(columns)]
        if not any(v not in (None, "") for v in r):
            break
        data.append(r + [None] * (len(columns) - len(r)))
    return columns, data


def read_table_analysis(path: Path) -> dict[str, dict]:
    """{TABLE: {"fields": str, "columns": [...], "rows": [...]}} -- one sheet per table: A1 "Table X", A2 the
    fields analysed, the grid from row 5."""
    wb = _load(path)
    try:
        out: dict[str, dict] = {}
        for ws in wb.worksheets:
            title = _text(ws.cell(1, 1).value)
            if not title.lower().startswith("table "):
                continue
            columns, rows = _grid(ws, 5)
            if columns:
                out[title[6:].strip().upper()] = {"fields": _text(ws.cell(2, 1).value), "columns": columns, "rows": rows}
        return out
    finally:
        wb.close()


_TEXT_COLUMNS = {"YEAR", "MONTH", "CATEGORY", "Y", "M", "CAT"}


def _numeric_columns(columns: list[str], rows: list[list]) -> list[list]:
    """SAP hands counts over as text ("6,930"): make the count columns real numbers (right-aligned, summable).
    Year, month and category columns stay as they are (a month is "07")."""
    rows = [list(r) for r in rows]
    for i, name in enumerate(columns):
        if name.upper() in _TEXT_COLUMNS:
            continue
        values = [r[i] for r in rows if r[i] not in (None, "")]
        if values and all(isinstance(v, (int, float)) or re.fullmatch(r"\d{1,3}(,\d{3})*|\d+", str(v).strip()) for v in values):
            for r in rows:
                if r[i] not in (None, "") and not isinstance(r[i], (int, float)):
                    r[i] = int(str(r[i]).replace(",", "").strip())
    return rows


def read_open_items(path: Path) -> dict[str, dict]:
    """{ARCHIVING OBJECT: {"columns", "rows"}}: one sheet per object (named after it), or -- in a file written
    by the earlier version -- one "Open item analysis" sheet with an "Archiving Object" column to split."""
    wb = _load(path)
    try:
        out: dict[str, dict] = {}
        for ws in wb.worksheets:
            if ws.title == "Summary":
                continue
            if _text(ws.cell(1, 1).value).lower() == "archiving object":          # old, single-sheet layout
                columns, rows = _grid(ws, 1)
                by_obj: dict[str, list] = {}
                for r in rows:
                    by_obj.setdefault(_text(r[0]).upper(), []).append(r[1:])
                for obj, obj_rows in by_obj.items():
                    keep = [i for i in range(len(columns) - 1)
                            if any(r[i] not in (None, "") for r in obj_rows)]     # drop columns other objects use
                    if obj and keep:
                        cols = [columns[i + 1] for i in keep]
                        out[obj] = {"columns": cols, "rows": _numeric_columns(cols, [[r[i] for i in keep] for r in obj_rows])}
            else:
                columns, rows = _grid(ws, 1)
                if columns and rows:
                    out[ws.title.strip().upper()] = {"columns": columns, "rows": _numeric_columns(columns, rows)}
        return out
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Styling helpers
# ---------------------------------------------------------------------------

_GREEN = PatternFill("solid", fgColor="86BC25")        # Deloitte green: text on it is black
_GREEN_LIGHT = PatternFill("solid", fgColor="EAF3D3")
_CHARCOAL = PatternFill("solid", fgColor="2B2B2B")
_GREY = PatternFill("solid", fgColor="E5E7EB")
_LINE = Side(style="thin", color="C9CCD1")
_GROUP_LINE = Side(style="medium", color="2B2B2B")
_BORDER = Border(left=_LINE, right=_LINE, top=_LINE, bottom=_LINE)
_MUTED = "6B7280"
_FONT = "Calibri"


def _lines(text: str, width: float) -> int:
    """How many lines *text* takes in a wrapped cell *width* characters wide."""
    per_line = max(1, int(width * 1.05))
    return sum(max(1, math.ceil(len(part) / per_line)) for part in str(text).split("\n"))


def _sheet_name(name: str, used: set) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", name)[:31] or "Object"
    out, n = base, 2
    while out.lower() in {u.lower() for u in used}:
        out = f"{base[:28]}_{n}"
        n += 1
    used.add(out)
    return out


def _to_bytes(wb) -> io.BytesIO:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Workbook 1: the overview
# ---------------------------------------------------------------------------

_OVERVIEW_COLUMNS = [
    ("Archiving Object", 22), ("Archiving Object Description", 38), ("Header Table", 20), ("Tables", 22),
    ("Archiving Conditions", 85), ("Archive Before", 26), ("Volume (GB)", 13), ("Volume (MB)", 14),
    ("Cumulative Size (GB)", 16), ("Cumulative Size (MB)", 17),
]
_HOUSEKEEPING_COLUMNS = [
    ("Housekeeping Program", 32), ("Tables", 24), ("Volume (GB)", 13), ("Volume (MB)", 14),
    ("Cumulative Size (GB)", 16), ("Cumulative Size (MB)", 17),
]


def conditions_text(conditions: list[str]) -> str:
    """One cell: the conditions, numbered, one per line."""
    return "\n".join(f"{n}. {c}" for n, c in enumerate(conditions, start=1))


def _write_grouped_sheet(
    ws, columns: list[tuple[str, int]], groups: list[dict], shared: set[int], numeric: set[int],
    wrap: set[int], top: set[int] = frozenset(),
) -> None:
    """A header row, then each group's rows. *groups* is [{"shared": {col: value}, "rows": [{col: value}]}] with
    1-based columns: the "shared" values (the columns in *shared*) are written once and merged down the group's
    rows; "rows" holds each row's own values. *numeric* columns are right-aligned with two decimals, *wrap*
    columns wrap, *top* columns are aligned to the top. The first column is bold."""
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "86BC25"
    ws.append([name for name, _ in columns])
    for col, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
        cell = ws.cell(1, col)
        cell.font = Font(name=_FONT, bold=True, color="FFFFFF")
        cell.fill = _CHARCOAL
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = _BORDER
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    widths = {i: w for i, (_, w) in enumerate(columns, start=1)}

    row_no = 2
    for g in groups:
        rows = g["rows"] or [{}]
        first, last = row_no, row_no + len(rows) - 1
        for i, values in enumerate(rows):
            r = first + i
            for col in range(1, len(columns) + 1):
                cell = ws.cell(r, col)
                # a group's shared value sits in its first row; the merge below spans the rest
                cell.value = g["shared"].get(col) if col in shared and i == 0 else values.get(col)
                cell.font = Font(name=_FONT, bold=(col == 1))
                cell.border = Border(left=_LINE, right=_LINE, bottom=_LINE, top=_GROUP_LINE if i == 0 else _LINE)
                cell.alignment = Alignment(
                    vertical="top" if col in top else "center", wrap_text=col in wrap,
                    horizontal="right" if col in numeric else "left",
                )
                if col in numeric:
                    cell.number_format = "#,##0.00"
        if len(rows) > 1:
            for col in shared:
                ws.merge_cells(start_row=first, start_column=col, end_row=last, end_column=col)

        # Rows do not grow to fit a merged, wrapped cell: give the group enough height for its longest text
        lines = [_lines(g["shared"][c], widths[c]) for c in shared & wrap if g["shared"].get(c)]
        per_row = max(18, math.ceil(max(lines + [1]) * 15 / len(rows)))
        for r in range(first, last + 1):
            ws.row_dimensions[r].height = per_row
        row_no = last + 1


def build_overview(
    groups: list[dict], headers: dict[str, str], conditions: dict[str, list[str]], archive_before: dict[str, str],
    housekeeping: list[dict],
) -> io.BytesIO:
    """The overview workbook: the "Overview" sheet (archiving objects) and the "Housekeeping Programs" sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Overview"
    _write_grouped_sheet(
        ws, _OVERVIEW_COLUMNS,
        [{
            "shared": {1: g["object"], 2: g["description"], 3: headers.get(g["object"], ""),
                       5: conditions_text(conditions.get(g["object"], [])), 6: archive_before.get(g["object"], ""),
                       9: g["cum_gb"], 10: g["cum_mb"]},
            "rows": [{4: t["table"], 7: t["gb"], 8: t["mb"]} for t in g["tables"]],
        } for g in groups],
        shared={1, 2, 3, 5, 6, 9, 10}, numeric={7, 8, 9, 10}, wrap={2, 4, 5, 6}, top={5, 6},
    )
    hk = wb.create_sheet("Housekeeping Programs")
    _write_grouped_sheet(
        hk, _HOUSEKEEPING_COLUMNS,
        [{
            "shared": {1: g["program"], 5: g["cum_gb"], 6: g["cum_mb"]},
            "rows": [{2: t["table"], 3: t["gb"], 4: t["mb"]} for t in g["tables"]],
        } for g in housekeeping],
        shared={1, 5, 6}, numeric={3, 4, 5, 6}, wrap={2},
    )
    return _to_bytes(wb)


# ---------------------------------------------------------------------------
# Workbook 2: one sheet per archiving object
# ---------------------------------------------------------------------------

def _write_block(ws, top: int, left: int, banner: str, columns: list[str], rows: list[list], note: str = "") -> int:
    """A banner, an optional italic note, a header row and the data; returns the number of columns used."""
    n = len(columns)
    widths = []
    for i, name in enumerate(columns):
        longest = max([len(str(name))] + [len(str(r[i])) for r in rows if r[i] is not None])
        widths.append(min(max(longest + 3, 12), 40))
    # a block narrower than its banner would cut the banner short
    while sum(widths) < len(banner) + 4:
        widths = [w + 2 for w in widths]
    for i, w in enumerate(widths):
        letter = get_column_letter(left + i)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, w)

    ws.cell(top, left, banner)
    for i in range(n):
        ws.cell(top, left + i).fill = _GREEN
    ws.cell(top, left).font = Font(name=_FONT, bold=True, size=12, color="000000")
    ws.cell(top, left).alignment = Alignment(vertical="center", indent=1)
    if n > 1:
        ws.merge_cells(start_row=top, start_column=left, end_row=top, end_column=left + n - 1)
    ws.row_dimensions[top].height = 24

    r = top + 1
    if note:
        ws.cell(r, left, note)
        ws.cell(r, left).font = Font(name=_FONT, italic=True, size=10, color=_MUTED)
        ws.cell(r, left).alignment = Alignment(vertical="top", wrap_text=True)
        if n > 1:
            ws.merge_cells(start_row=r, start_column=left, end_row=r, end_column=left + n - 1)
        ws.row_dimensions[r].height = max(16, _lines(note, sum(widths)) * 14)
    r = top + 2  # the header row is on the same line for every block, so blocks line up side by side
    for i, name in enumerate(columns):
        cell = ws.cell(r, left + i, name)
        cell.font = Font(name=_FONT, bold=True)
        cell.fill = _GREY
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 22)
    for row in rows:
        r += 1
        is_total = _text(row[0]).lower() == "total"
        for i, value in enumerate(row):
            cell = ws.cell(r, left + i, value)
            cell.border = _BORDER
            cell.font = Font(name=_FONT, bold=is_total)
            cell.alignment = Alignment(horizontal="right" if isinstance(value, (int, float)) else "left", vertical="center")
            if is_total:
                cell.fill = _GREEN_LIGHT
            if isinstance(value, int) and not isinstance(value, bool):
                cell.number_format = "#,##0"
    return n


def build_details(
    groups: list[dict], headers: dict[str, str],
    table_analysis: dict[str, dict], open_items: dict[str, dict],
) -> tuple[io.BytesIO, dict[str, dict]]:
    """The per-object workbook, and for each object which parts it got: {object: {"table_analysis", "open_items"}}.
    An object with neither has no sheet (nothing to show)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used: set = set()
    parts: dict[str, dict] = {}

    for g in groups:
        obj = g["object"]
        header_table = headers.get(obj, "")
        ta = table_analysis.get(header_table) if header_table else None
        oi = open_items.get(obj)
        parts[obj] = {"table_analysis": bool(ta), "open_items": bool(oi)}
        if not ta and not oi:
            continue

        ws = wb.create_sheet(_sheet_name(obj, used))
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = "86BC25"
        ws.cell(1, 1, f"Archiving object: {obj}").font = Font(name=_FONT, bold=True, size=16)
        ws.row_dimensions[1].height = 28
        ws.cell(2, 1, g["description"]).font = Font(name=_FONT, size=11, color=_MUTED)
        if header_table:
            ws.cell(3, 1, f"Header table: {header_table}").font = Font(name=_FONT, bold=True)
        if header_table and not ta:
            ws.cell(4, 1, f"No table analysis is available for the header table {header_table}.").font = \
                Font(name=_FONT, italic=True, size=10, color=_MUTED)
        elif not header_table:
            ws.cell(4, 1, "No header table was found for this archiving object, so there is no table analysis.").font = \
                Font(name=_FONT, italic=True, size=10, color=_MUTED)

        left = 1
        if ta:
            left += _write_block(ws, 6, left, f"Table analysis — {header_table}", ta["columns"], ta["rows"], ta["fields"]) + 1
        if oi:
            _write_block(ws, 6, left, "Open item analysis", oi["columns"], oi["rows"])
    if not wb.worksheets:
        wb.create_sheet("No data").cell(1, 1, "None of the archiving objects has a table analysis or an open item analysis yet.")
    return _to_bytes(wb), parts


# ---------------------------------------------------------------------------
# The whole run
# ---------------------------------------------------------------------------

def generate(save: Callable[[str, str, io.BytesIO], str]) -> dict:
    """Read the sources, build both workbooks and save them with *save(task, filename, buffer) -> path*.
    Returns {"files": [{"name", "path"}], "objects": [...], "warnings": [...]}."""
    sources = resolve_sources()
    warnings: list[str] = [s["warning"] for s in sources.values() if s["warning"]]

    objects_path = sources["objects"]["path"]
    if objects_path is None:
        raise ReportError("There are no archiving objects yet. Run Find Archiving Objects for Tables first "
                          "(it saves archiving_objects_scored.xlsx).")
    groups = read_groups(objects_path)
    if not groups:
        raise ReportError(f"{objects_path.name} has no archiving objects on its \"Grouped by Object\" sheet.")

    def optional(key: str, reader, empty, why: str):
        path = sources[key]["path"]
        if path is None:
            warnings.append(f"No {sources[key]['label'].lower()} found: {why}")
            return empty
        return reader(path)

    headers = optional("headers", read_header_tables, {}, "the Header Table column is blank and there are no object sheets with table analysis.")
    conditions = optional("conditions", read_conditions, {}, "the Archiving Conditions column is blank.")
    archive_before = optional("conditions", read_archive_before, {}, "the Archive Before column is blank.")
    table_analysis = optional("table_analysis", read_table_analysis, {}, "the object sheets have no table analysis.")
    open_items = optional("open_items", read_open_items, {}, "the object sheets have no open item analysis.")

    housekeeping = read_housekeeping_groups(objects_path)
    overview = build_overview(groups, headers, conditions, archive_before, housekeeping)
    details, parts = build_details(groups, headers, table_analysis, open_items)
    folder_task = "FINAL_REPORT"
    files = [
        {"name": OVERVIEW_NAME, "path": save(folder_task, OVERVIEW_NAME, overview)},
        {"name": DETAILS_NAME, "path": save(folder_task, DETAILS_NAME, details)},
    ]

    rows = [{
        "object": g["object"], "description": g["description"], "header_table": headers.get(g["object"], ""),
        "tables": len(g["tables"]), "conditions": len(conditions.get(g["object"], [])),
        "table_analysis": parts[g["object"]]["table_analysis"], "open_items": parts[g["object"]]["open_items"],
    } for g in groups]
    with_sheet = sum(1 for p in parts.values() if p["table_analysis"] or p["open_items"])
    if with_sheet < len(groups):
        warnings.append(f"{len(groups) - with_sheet} of {len(groups)} archiving objects have no table analysis or open item "
                        "analysis yet, so they have no sheet in the object details workbook.")
    return {"files": files, "objects": rows, "warnings": warnings,
            "housekeeping": {"programs": len(housekeeping), "tables": sum(len(g["tables"]) for g in housekeeping)}}
