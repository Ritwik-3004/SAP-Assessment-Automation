"""
Table analysis task: TAANA ad hoc analyses for the header tables of a header-table sheet.

For each table, in order:
  1. read the fields TAANA offers (transactions/taana_analysis.py) and the date/year-like ones
     (ABAP dictionary types, looked up once for all tables);
  2. if the table is listed in resources/Fields for TAANA.xlsx, use the fields given there and skip the
     questions; otherwise ask the user which DATE/YEAR/MONTH fields to analyse (and whether to group dates
     by year);
  3. (only for tables not in the sheet) ask whether other fields (company code, document type, ...) should
     be added; if yes, ask which;
  4. create the ad hoc variant and schedule the analysis in the background (start immediately) -- then
     go straight on to the next table's questions, without waiting for the analysis;
  5. meanwhile a second track watches the scheduled jobs and, as each completes, copies its result grid
     into  output/Table analysis.xlsx.

When the whole run is done, any analysed table can be re-run with additional fields (redo()): the user picks
the extra fields, the analysis is repeated with the earlier fields plus the new ones, and the table's result
in the workbook is replaced.

The job runs on background threads and talks to the UI through a small state object: the UI polls
snapshot() and answers the current question through answer(). Questions are asked between SAP steps, so
SAP is not held while the user thinks.
"""

import io
import logging
import re
import threading
import time
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from config import OUTPUT_DIR
import transactions.db02 as db02
import transactions.taana_analysis as taana

logger = logging.getLogger(__name__)

OUTPUT_NAME = "Table analysis.xlsx"
POLL_SECONDS = 10.0          # between checks of a running background job
FIRST_POLL_SECONDS = 4.0
MAX_WAIT_SECONDS = 4 * 3600  # give up on one table's job after this long (the user can also Skip)
MAX_CHECK_ERRORS = 3

# Fields worth suggesting first when the user adds "other" fields.
SUGGESTED_FIELDS = {
    "BUKRS", "BLART", "MONAT", "KOKRS", "WERKS", "VKORG", "VTWEG", "SPART", "EKORG", "EKGRP",
    "LGORT", "OBJECTCLAS", "TCODE", "USNAM", "ERNAM", "BSTAT", "VBTYP", "AUART", "BSART", "LAND1",
}
EXCLUDED_OTHER = {"MANDT"}

_TABLE_NAME = re.compile(r"^[A-Z0-9_/]{2,30}$")
_YEAR_NAME = re.compile(r"(GJAHR|JAHR|YEAR|STJAH|FYEAR)")
_YEAR_LABEL = re.compile(r"\b(year|yr|jahr)\b", re.IGNORECASE)
_MONTH_NAME = re.compile(r"(MONAT|MONTH|MON$|^MON|_MON|SPBUP|POPER|BUPER|PERIO|PERID|PERBL|PERAB|PERBI)")
_MONTH_LABEL = re.compile(r"\b(month|monat|period|periode)\b", re.IGNORECASE)


FIELDS_SHEET = Path(__file__).parent / "resources" / "Fields for TAANA.xlsx"
_sheet_cache: dict = {"mtime": None, "data": None}


def load_fields_sheet() -> dict:
    """The fields to analyse per table, from resources/Fields for TAANA.xlsx (re-read when the file changes).

    The sheet needs a "Header Table" column (a "Table Name"/"Table" column also works) and one or more
    columns whose title contains "field" (e.g. "Primary Date Field"); a cell may hold several fields
    separated by commas, semicolons or line breaks. Rows for the same table are combined.

    Returns {"tables": {TABLE: [FIELD, ...]}, "problem": str}; "problem" is non-empty (and "tables" empty)
    if the sheet is missing or unusable -- the task then just asks the user, as it did before the sheet."""
    try:
        mtime = FIELDS_SHEET.stat().st_mtime
    except OSError:
        return {"tables": {}, "problem": f"The fields sheet was not found ({FIELDS_SHEET.name} in backend/resources)."}
    if _sheet_cache["mtime"] == mtime and _sheet_cache["data"] is not None:
        return _sheet_cache["data"]

    data: dict = {"tables": {}, "problem": ""}
    try:
        wb = openpyxl.load_workbook(FIELDS_SHEET, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None) or ()
            table_col = next((i for i, h in enumerate(header) if _norm(h) in ("headertable", "tablename", "table")), None)
            field_cols = [i for i, h in enumerate(header) if "field" in _norm(h)]
            if table_col is None or not field_cols:
                data["problem"] = "The fields sheet needs a 'Header Table' column and a column with 'Field' in its title."
            else:
                for row in rows:
                    table = str(row[table_col] if table_col < len(row) and row[table_col] is not None else "").strip().upper()
                    if not table or not _TABLE_NAME.match(table):
                        continue
                    names = data["tables"].setdefault(table, [])
                    for col in field_cols:
                        cell = row[col] if col < len(row) else None
                        for part in re.split(r"[,;\n]+", str(cell or "")):
                            name = part.strip().upper()
                            if name and name not in names:
                                names.append(name)
                data["tables"] = {t: f for t, f in data["tables"].items() if f}
        finally:
            wb.close()
    except Exception as exc:
        logger.warning("Could not read the fields sheet: %s", exc)
        data = {"tables": {}, "problem": f"The fields sheet could not be read ({exc})."}
    _sheet_cache.update(mtime=mtime, data=data)
    return data


class _Skipped(Exception):
    pass


class _Cancelled(Exception):
    pass


# ---------------------------------------------------------------------------
# Input: the header tables
# ---------------------------------------------------------------------------

def _norm(text) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def parse_tables(contents: bytes) -> list[str]:
    """Distinct header tables from a header-table workbook, in sheet order.

    Uses the LAST sheet that has a "Header Table" column (for the reference-reviewed workbook that is the
    final list after the user's overrides). Falls back to a "Table Name"/"Table" column. Blank cells and
    names that are not valid table names are skipped. Raises ValueError with a readable message."""
    workbook = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    try:
        chosen: Optional[tuple] = None
        for wanted in (("headertable",), ("tablename", "table")):
            for ws in workbook.worksheets:
                header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
                if not header:
                    continue
                for idx, cell in enumerate(header):
                    if _norm(cell) in wanted:
                        chosen = (ws, idx)
                        break
            if chosen:
                break
        if not chosen:
            raise ValueError(
                "No 'Header Table' column found. Use the file saved by Find Header Tables "
                "(header_tables.xlsx / header_tables_with_reference.xlsx)."
            )
        ws, idx = chosen
        tables: list[str] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            value = str(row[idx] if idx < len(row) and row[idx] is not None else "").strip().upper()
            if value and _TABLE_NAME.match(value) and value not in tables:
                tables.append(value)
        if not tables:
            raise ValueError("The 'Header Table' column has no table names.")
        return tables
    finally:
        workbook.close()


# ---------------------------------------------------------------------------
# Which fields are dates / years
# ---------------------------------------------------------------------------

def _year_like(name: str, label: str) -> bool:
    return bool(_YEAR_NAME.search(name) or _YEAR_LABEL.search(label or ""))


def _month_like(name: str, label: str) -> bool:
    return bool(_MONTH_NAME.search(name) or _MONTH_LABEL.search(label or ""))


def classify_fields(fields: list[dict], types: dict, types_known: bool) -> tuple[list[dict], list[dict]]:
    """Split TAANA's fields into (date/year/month fields, other fields).

    *types* is {FIELD: {"type", "length"}} from the ABAP dictionary for the date-like candidates of this
    table. Kinds: "date" (DATS), "period" (YYYYMM: ACCP, or a 6-character field named like a period),
    "year" (4-digit NUMC named/labelled as a year), "month" (a 2/3-character month or posting-period field
    such as MONAT or POPER). Dates and periods can be grouped by year; a year or month field already is
    one. If the dictionary lookup failed (*types_known* False) names and labels are used instead, so the
    task still works."""
    date_fields: list[dict] = []
    other_fields: list[dict] = []
    for f in fields:
        name, label = f["name"].upper(), f.get("label", "")
        kind = None
        info = types.get(name)
        if info:
            if info["type"] == "DATS":
                kind = "date"
            elif info["type"] == "ACCP":
                kind = "period"
            elif info["type"] in ("NUMC", "CHAR") and info["length"] == 4 and _year_like(name, label):
                kind = "year"
            elif info["type"] in ("NUMC", "CHAR") and info["length"] == 6 and _month_like(name, label):
                kind = "period"
            elif info["type"] in ("NUMC", "CHAR") and info["length"] in (2, 3) and _month_like(name, label):
                kind = "month"
        elif not types_known:
            lowered = label.lower()
            if _year_like(name, label):
                kind = "year"
            elif _month_like(name, label):
                kind = "month"
            elif name.endswith("DAT") or "DATE" in name or "DATUM" in name or re.search(r"\bdate\b", lowered):
                kind = "date"
        if kind:
            date_fields.append({"name": name, "label": label, "kind": kind})
        elif name not in EXCLUDED_OTHER:
            other_fields.append({"name": name, "label": label, "suggested": name in SUGGESTED_FIELDS})
    return date_fields, other_fields


# ---------------------------------------------------------------------------
# The job
# ---------------------------------------------------------------------------

def _clock(seconds: float) -> str:
    seconds = int(seconds)
    return f"{seconds // 3600}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


class TableAnalysisJob:
    """Two tracks run side by side:

    * the QUESTION track (``_run``) goes through the tables in order, asks the user about the fields and
      schedules each table's SAP background job straight away, then moves on to the next table;
    * the COLLECTOR track (``_collect``) watches the scheduled jobs, and as each one completes copies its
      result into the workbook.
    So the user answers all the questions without waiting for any analysis. Every SAP call goes through
    the single SAP worker thread, so the two tracks never touch the SAP screen at the same moment.
    """

    def __init__(self):
        self._lock = threading.RLock()
        self._answered = threading.Event()
        self._interactive_done = threading.Event()
        self._group_by_year = True
        self._types: dict = {}
        self._types_known = False
        self._reset()

    def _reset(self) -> None:
        self._state = {
            "status": "idle",         # idle | running | waiting (for the user) | done | error
            "mode": "run",            # "run", or "redo" while one table is re-run with additional fields
            "total": 0,
            "asked": 0,               # tables whose questions are answered (or that were skipped)
            "completed": 0,           # tables finished: analysed, skipped or failed
            "message": None,
            "current": None,          # the table being asked about: {"table", "step"}
            "prompt": None,           # the open question, if status == "waiting"
            "interactive_done": False,
            "records": [],            # one entry per table handled so far, in input order
            "running": [],            # tables whose SAP job is still running
            "output_path": None,
            "warnings": [],
        }
        self._records: dict[str, dict] = {}
        self._index: dict[str, int] = {}          # table -> position in the input
        self._pending: dict[str, dict] = {}       # scheduled jobs not yet read
        self._results: list[dict] = []            # full result grids, for the workbook
        self._answer: Optional[dict] = None
        self._prompt_seq = 0
        self._cancel = False
        self._skip = False
        self._stop_waiting: set[str] = set()
        self._answered.clear()
        self._interactive_done.clear()

    # -- public API -----------------------------------------------------------

    def snapshot(self) -> dict:
        with self._lock:
            state = dict(self._state)
            state["records"] = [
                dict(self._records[t]) for t in sorted(self._records, key=lambda t: self._index.get(t, 0))
            ]
            state["running"] = sorted(self._pending, key=lambda t: self._index.get(t, 0))
            state["warnings"] = list(self._state["warnings"])
        sheet = load_fields_sheet()
        state["fields_sheet"] = {"tables": len(sheet["tables"]), "problem": sheet["problem"]}
        return state

    def is_active(self) -> bool:
        with self._lock:
            return self._state["status"] in ("running", "waiting")

    def start(self, tables: list[str], group_by_year: bool = True) -> None:
        """*group_by_year*: for tables taken from the fields sheet, group date fields by year."""
        with self._lock:
            if self._state["status"] in ("running", "waiting"):
                raise RuntimeError("A table analysis is already running.")
            self._reset()
            self._group_by_year = group_by_year
            self._state.update(status="running", total=len(tables), message="Starting…")
        threading.Thread(target=self._run, args=(tables,), daemon=True, name="table-analysis").start()

    def redo(self, table: str) -> None:
        """Re-run an analysed table with additional fields (only once the whole run is done): the user is asked
        which extra fields to add, the earlier fields are kept, and the table's result is replaced."""
        table = table.strip().upper()
        with self._lock:
            if self._state["status"] != "done":
                raise RuntimeError("Wait until the analysis has finished before re-running a table.")
            record = self._records.get(table)
            if not record or record["state"] != "done" or not record.get("specs"):
                raise RuntimeError(f"{table} has no finished analysis to re-run.")
            self._cancel = False
            self._skip = False
            self._answered.clear()
            self._interactive_done.clear()
            self._state.update(
                status="running", mode="redo", interactive_done=False, current=None,
                message=f"Re-running {table} with additional fields…",
            )
        threading.Thread(target=self._run_redo, args=(table,), daemon=True, name="table-analysis-redo").start()

    def answer(self, prompt_id: int, answer: dict) -> bool:
        with self._lock:
            prompt = self._state["prompt"]
            if not prompt or prompt["id"] != prompt_id:
                return False
            self._answer = answer
        self._answered.set()
        return True

    def skip(self) -> None:
        """Skip the table the questions are being asked about."""
        with self._lock:
            if not self._interactive_done.is_set():
                self._skip = True

    def stop_waiting(self, table: str) -> bool:
        """Stop waiting for a table whose SAP job is running (the job itself keeps running in SAP)."""
        with self._lock:
            if table not in self._pending:
                return False
            self._stop_waiting.add(table)
            return True

    def cancel(self) -> None:
        """Stop asking and collecting. Finished tables stay in the workbook; jobs already running keep
        running in SAP."""
        with self._lock:
            self._cancel = True

    # -- plumbing -------------------------------------------------------------

    def _set(self, **changes) -> None:
        with self._lock:
            self._state.update(changes)

    def _step(self, table: str, step: str, message: Optional[str] = None) -> None:
        self._set(current={"table": table, "step": step}, message=message or step)

    def _check_flags(self) -> None:
        with self._lock:
            if self._cancel:
                raise _Cancelled()
            if self._skip:
                self._skip = False
                raise _Skipped("Skipped by the user.")

    def _ask(self, kind: str, payload: dict) -> dict:
        """Put a question to the user and block (SAP is not held) until they answer."""
        with self._lock:
            self._prompt_seq += 1
            self._answer = None
            self._answered.clear()
            self._state["prompt"] = {"id": self._prompt_seq, "kind": kind, **payload}
            self._state["status"] = "waiting"
        try:
            while not self._answered.wait(0.4):
                self._check_flags()
            with self._lock:
                return self._answer or {}
        finally:
            with self._lock:
                self._state["prompt"] = None
                if self._state["status"] == "waiting":
                    self._state["status"] = "running"

    def _finish(self, table: str, record: dict, result: Optional[dict] = None) -> None:
        """A table is finished (analysed, skipped or failed): record it and refresh the workbook."""
        with self._lock:
            self._records[table] = record
            self._pending.pop(table, None)
            self._stop_waiting.discard(table)
            self._state["completed"] = sum(1 for r in self._records.values() if r["state"] != "running")
            if result is not None:
                self._results = [r for r in self._results if r["table"] != table]   # a re-run replaces the old one
                self._results.append({**result, "index": self._index.get(table, 0)})
        try:
            self._write_workbook()
        except Exception as exc:
            logger.exception("Could not write the Table analysis workbook")
            with self._lock:
                self._state["warnings"].append(f"Could not save the Excel file: {exc}")

    # -- question track -------------------------------------------------------

    def _run(self, tables: list[str]) -> None:
        collector: Optional[threading.Thread] = None
        try:
            self._set(message="Reading the date/year/month fields of all tables from the ABAP dictionary…")
            lookup = db02.run_get_field_types(tables)
            types_known = lookup.get("status") == "ok"
            all_types = lookup.get("types", {}) if types_known else {}
            self._types, self._types_known = all_types, types_known
            sheet = load_fields_sheet()
            if sheet["problem"]:
                with self._lock:
                    self._state["warnings"].append(f"{sheet['problem']} You will be asked for every table's fields.")
            if not types_known:
                with self._lock:
                    self._state["warnings"].append(
                        "Could not read field types from the ABAP dictionary "
                        f"({lookup.get('message', 'unknown error')}); date/year/month fields are guessed from their names."
                    )

            collector = threading.Thread(target=self._collect, daemon=True, name="table-analysis-collector")
            collector.start()

            for number, table in enumerate(tables, start=1):
                with self._lock:
                    if self._cancel:
                        break
                    self._skip = False
                    self._index[table] = number
                try:
                    entry = self._submit(table, all_types.get(table, {}), types_known, sheet["tables"].get(table))
                except _Skipped as skipped:
                    self._finish(table, {"table": table, "state": "skipped", "fields": [], "rows": 0, "note": str(skipped)})
                except _Cancelled:
                    break
                except Exception as exc:
                    logger.exception("Table analysis failed for %s", table)
                    self._finish(table, {"table": table, "state": "failed", "fields": [], "rows": 0, "note": str(exc)})
                else:
                    with self._lock:
                        self._pending[table] = entry
                        self._records[table] = {
                            "table": table, "state": "running", "fields": entry["fields"], "rows": 0,
                            "source": entry["source"], "specs": entry["specs"],
                            "note": "Job scheduled in SAP; waiting for it to finish."
                            + (f" {entry['sheet_note']}" if entry.get("sheet_note") else ""),
                        }
                with self._lock:
                    self._state["asked"] = number
        except Exception as exc:
            logger.exception("Table analysis job failed")
            with self._lock:
                self._state["status"] = "error"
                self._state["message"] = str(exc)
            self._interactive_done.set()
            return

        with self._lock:
            self._state["current"] = None
            self._state["interactive_done"] = True
            self._state["message"] = "All questions answered."
        self._interactive_done.set()
        if collector is not None:
            collector.join()
        with self._lock:
            left = len(self._pending)
            self._state["status"] = "done"
            self._state["message"] = (
                f"Cancelled. {left} job(s) already scheduled keep running in SAP." if self._cancel and left
                else "Cancelled." if self._cancel
                else "Finished."
            )

    def _submit(self, table: str, types: dict, types_known: bool, sheet_fields: Optional[list[str]] = None) -> dict:
        """Work out the fields for *table* and schedule its SAP job. Returns the pending-job entry.

        If the fields sheet lists the table, those fields are used and the user is NOT asked. The user is asked
        only when the table is not in the sheet, or none of its fields exists in TAANA for this table."""
        self._step(table, "Reading the table's fields in TAANA…")
        listed = taana.list_fields(table)
        if listed["status"] != "ok":
            raise RuntimeError(listed["message"])
        date_fields, other_fields = classify_fields(listed["fields"], types, types_known)

        specs: list[dict] = []
        source = "you"
        sheet_note = ""
        sheet_problem = ""
        if sheet_fields:
            offered = {f["name"].upper() for f in listed["fields"]}
            kinds = {f["name"]: f["kind"] for f in date_fields}
            usable = [n for n in sheet_fields if n in offered]
            missing = [n for n in sheet_fields if n not in offered]
            if usable:
                specs = [
                    {"name": n, "year": self._group_by_year and kinds.get(n) in ("date", "period")} for n in usable
                ]
                source = "sheet"
                sheet_note = "Fields taken from the Fields for TAANA sheet."
                if missing:
                    sheet_note += f" Not offered by TAANA for this table and left out: {', '.join(missing)}."
                self._step(table, f"Using the fields from the sheet: {', '.join(usable)}")
            else:
                sheet_problem = (
                    f"The Fields for TAANA sheet lists {', '.join(sheet_fields)} for {table}, but TAANA does not "
                    "offer any of them for this table, so please choose the fields."
                )

        if not specs:
            if date_fields:
                self._step(table, "Waiting for you to choose the date/year/month fields", "Choose the date/year/month fields.")
                answer = self._ask("date_fields", {
                    "table": table,
                    "fields": date_fields,
                    "can_group_by_year": any(f["kind"] in ("date", "period") for f in date_fields),
                    "note": sheet_problem,
                })
                if answer.get("skip"):
                    raise _Skipped("Skipped by the user.")
                by_name = {f["name"]: f for f in date_fields}
                group = bool(answer.get("group_by_year"))
                for name in answer.get("selected", []):
                    f = by_name.get(str(name).upper())
                    if f:
                        specs.append({"name": f["name"], "year": group and f["kind"] in ("date", "period")})

            self._step(table, "Waiting for you: add other fields?", "Add other fields?")
            more = self._ask("more_fields", {
                "table": table,
                "no_date_fields": not date_fields,
                "selected": [s["name"] for s in specs],
                "has_other_fields": bool(other_fields),
                "note": sheet_problem if not date_fields else "",
            })
            if more.get("skip"):
                raise _Skipped("Skipped by the user.")
            if more.get("add") and other_fields:
                self._step(table, "Waiting for you to choose the other fields", "Choose the other fields.")
                picked = self._ask("other_fields", {"table": table, "fields": other_fields})
                if picked.get("skip"):
                    raise _Skipped("Skipped by the user.")
                valid = {f["name"] for f in other_fields}
                specs += [{"name": str(n).upper(), "year": False} for n in picked.get("selected", []) if str(n).upper() in valid]

        if not specs:
            raise _Skipped("No fields were selected.")
        return self._schedule(table, specs, source, sheet_note)

    def _schedule(self, table: str, specs: list[dict], source: str, sheet_note: str = "", previous: Optional[dict] = None) -> dict:
        self._step(table, "Creating the ad hoc variant and starting the background job in SAP…")
        started = taana.start_analysis(table, specs)
        if started["status"] != "ok":
            raise RuntimeError(started["message"])
        now = time.monotonic()
        return {
            "table": table,
            "before": started["before"],
            "fields": [s["name"] + ("(year)" if s["year"] else "") for s in specs],
            "specs": specs,
            "source": source,
            "sheet_note": sheet_note,
            "previous": previous,
            "submitted": now,
            "next_check": now + FIRST_POLL_SECONDS,
            "errors": 0,
        }

    # -- re-run with additional fields ------------------------------------------

    def _submit_redo(self, table: str) -> dict:
        with self._lock:
            previous = dict(self._records[table])
        old_specs = [dict(s) for s in previous["specs"]]
        self._step(table, "Reading the table's fields in TAANA…", f"Re-running {table}: reading its fields…")
        listed = taana.list_fields(table)
        if listed["status"] != "ok":
            raise RuntimeError(listed["message"])
        date_fields, other_fields = classify_fields(listed["fields"], self._types.get(table, {}), self._types_known)
        used = {s["name"] for s in old_specs}
        offered = [f for f in date_fields if f["name"] not in used] + sorted(
            (f for f in other_fields if f["name"] not in used), key=lambda f: not f["suggested"]
        )
        if not offered:
            raise RuntimeError("There are no other fields left to add.")
        self._step(table, "Waiting for you to choose the additional fields", "Choose the additional fields.")
        answer = self._ask("redo_fields", {
            "table": table,
            "previous": [s["name"] + ("(year)" if s["year"] else "") for s in old_specs],
            "fields": offered,
            "can_group_by_year": any(f.get("kind") in ("date", "period") for f in offered),
        })
        if answer.get("skip"):
            raise _Skipped("Re-run cancelled.")
        by_name = {f["name"]: f for f in offered}
        group = bool(answer.get("group_by_year"))
        added = []
        for name in answer.get("selected", []):
            f = by_name.get(str(name).upper())
            if f:
                added.append({"name": f["name"], "year": group and f.get("kind") in ("date", "period")})
        if not added:
            raise _Skipped("No additional fields were selected; the previous result was kept.")
        entry = self._schedule(table, old_specs + added, "re-run", "Re-run with additional fields: "
                               + ", ".join(a["name"] for a in added) + ".", previous=previous)
        return entry

    def _run_redo(self, table: str) -> None:
        collector = threading.Thread(target=self._collect, daemon=True, name="table-analysis-collector")
        collector.start()
        note = ""
        try:
            entry = self._submit_redo(table)
        except (_Skipped, _Cancelled) as exc:
            note = str(exc) or "Re-run cancelled."
        except Exception as exc:
            logger.exception("Re-run failed for %s", table)
            note = f"Re-run failed: {exc}"
        else:
            with self._lock:
                self._pending[table] = entry
                self._records[table] = {
                    **self._records[table], "state": "running", "fields": entry["fields"], "specs": entry["specs"],
                    "source": "re-run", "note": f"Re-running in SAP. {entry['sheet_note']}",
                }
        if note:
            with self._lock:
                rec = self._records.get(table)
                if rec is not None:
                    rec["note"] = f"{note} The previous result was kept."
        with self._lock:
            self._state["current"] = None
            self._state["interactive_done"] = True
        self._interactive_done.set()
        collector.join()
        with self._lock:
            self._state.update(status="done", mode="run", message="Re-run finished." if not note else note)

    def _unsuccessful(self, entry: dict, state: str, note: str) -> dict:
        """The record for a job that did not give a result. A re-run that fails keeps the previous result."""
        previous = entry.get("previous")
        if previous:
            return {**previous, "note": f"{note} The previous result was kept."}
        return {
            "table": entry["table"], "state": state, "fields": entry["fields"], "specs": entry["specs"],
            "source": entry["source"], "rows": 0, "note": note,
        }

    # -- collector track ------------------------------------------------------

    def _collect(self) -> None:
        """Check the scheduled jobs and write each result as its job completes, until the question
        track is done and nothing is left to wait for (or the run is cancelled)."""
        while True:
            with self._lock:
                if self._cancel:
                    return
                pending = list(self._pending.values())
                stopped = set(self._stop_waiting)
            if not pending:
                if self._interactive_done.is_set():
                    return
                time.sleep(0.5)
                continue

            for entry in [p for p in pending if p["table"] in stopped]:
                self._finish(entry["table"], self._unsuccessful(
                    entry, "skipped", "Stopped waiting; the job keeps running in SAP."
                ))

            now = time.monotonic()
            due = [p for p in pending if p["table"] not in stopped and p["next_check"] <= now]
            if not due:
                time.sleep(0.5)
                continue

            results = taana.check_analyses([{"table": p["table"], "before": p["before"]} for p in due])
            now = time.monotonic()
            for entry in due:
                table = entry["table"]
                result = results.get(table) or {"state": "error", "status": "No answer from TAANA."}
                state = result.get("state")
                elapsed = now - entry["submitted"]
                if state == "completed":
                    self._finish(
                        table,
                        {"table": table, "state": "done", "fields": entry["fields"], "rows": len(result["rows"]),
                         "specs": entry["specs"], "source": entry["source"],
                         "note": f"Analysis {result.get('status', 'completed')}."
                         + (f" {entry['sheet_note']}" if entry.get("sheet_note") else "")},
                        {"table": table, "fields": entry["fields"], "started": result.get("started", ""),
                         "status": result.get("status", ""), "columns": result["columns"],
                         "column_ids": result["column_ids"], "rows": result["rows"]},
                    )
                    continue
                if state == "failed":
                    self._finish(table, self._unsuccessful(
                        entry, "failed", f"The SAP analysis job ended with status '{result.get('status')}'."
                    ))
                    continue
                entry["errors"] = entry["errors"] + 1 if state == "error" else 0
                if entry["errors"] >= MAX_CHECK_ERRORS or elapsed > MAX_WAIT_SECONDS:
                    reason = (
                        f"Could not read the analysis in TAANA: {result.get('status')}"
                        if entry["errors"] >= MAX_CHECK_ERRORS
                        else f"The analysis did not finish within {_clock(MAX_WAIT_SECONDS)}."
                    )
                    self._finish(table, self._unsuccessful(entry, "failed", reason))
                    continue
                entry["next_check"] = now + POLL_SECONDS
                with self._lock:
                    if table in self._records:
                        self._records[table]["note"] = f"Running in SAP ({_clock(elapsed)}): {result.get('status', '')}"

            with self._lock:
                waiting = len(self._pending)
                if self._interactive_done.is_set() and waiting:
                    self._state["message"] = f"All questions answered. Waiting for {waiting} SAP job{'s' if waiting != 1 else ''}…"

    # -- output ---------------------------------------------------------------

    def _write_workbook(self) -> None:
        wb = openpyxl.Workbook()
        summary = wb.active
        summary.title = "Summary"
        summary.append(["Table", "Result", "Fields analysed", "Fields chosen by", "Result rows", "Analysis started", "Note"])
        with self._lock:
            results = sorted(self._results, key=lambda r: r["index"])
            records = [dict(self._records[t]) for t in sorted(self._records, key=lambda t: self._index.get(t, 0))]
        started_by_table = {r["table"]: r["started"] for r in results}
        for r in records:
            summary.append([
                r["table"], r["state"].capitalize(), ", ".join(r["fields"]),
                {"sheet": "Fields for TAANA sheet", "re-run": "You (re-run with additional fields)"}.get(r.get("source", ""), "You"),
                r["rows"], started_by_table.get(r["table"], ""), r["note"],
            ])
        self._style_header(summary, 1)
        self._autosize(summary)

        used = {"Summary"}
        for res in results:
            ws = wb.create_sheet(self._sheet_name(res["table"], used))
            ws.append([f"Table {res['table']}"])
            ws["A1"].font = Font(bold=True, size=13)
            ws.append([f"Fields analysed: {', '.join(res['fields'])}"])
            ws.append([f"Analysis started: {res['started']}    Status: {res['status']}"])
            ws.append([])
            ws.append(res["columns"])
            self._style_header(ws, 5)
            count_idx = res["column_ids"].index("_ENTRY_CNT") if "_ENTRY_CNT" in res["column_ids"] else len(res["columns"]) - 1
            for row in res["rows"]:
                out = list(row)
                digits = "".join(ch for ch in str(out[count_idx]) if ch.isdigit())
                if digits:
                    out[count_idx] = int(digits)
                # SAP ends the grid with a total line whose field cells are blank
                if digits and not any(str(v).strip() for i, v in enumerate(out) if i != count_idx):
                    out[0] = "Total"
                ws.append(out)
            ws.freeze_panes = "A6"
            self._autosize(ws, skip_rows=4)

        path = self._save(wb)
        self._set(output_path=str(path))

    @staticmethod
    def _sheet_name(table: str, used: set) -> str:
        base = re.sub(r"[\\/*?:\[\]]", "_", table)[:31] or "Table"
        name, n = base, 2
        while name in used:
            name = f"{base[:28]}_{n}"
            n += 1
        used.add(name)
        return name

    @staticmethod
    def _style_header(ws, row: int) -> None:
        for cell in ws[row]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center")

    @staticmethod
    def _autosize(ws, skip_rows: int = 0) -> None:
        widths: dict[int, int] = {}
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if r_idx <= skip_rows:
                continue
            for c_idx, value in enumerate(row, start=1):
                widths[c_idx] = max(widths.get(c_idx, 8), min(len(str(value)) + 2 if value is not None else 0, 60))
        for c_idx, width in widths.items():
            ws.column_dimensions[get_column_letter(c_idx)].width = width

    @staticmethod
    def _save(wb) -> Path:
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        candidates = [OUTPUT_DIR / OUTPUT_NAME] + [
            OUTPUT_DIR / f"{Path(OUTPUT_NAME).stem} ({n}){Path(OUTPUT_NAME).suffix}" for n in range(2, 10)
        ]
        for path in candidates:
            try:
                wb.save(path)
                return path
            except PermissionError:  # open in Excel
                continue
        raise PermissionError("Could not write the Table analysis workbook (is it open in Excel?).")


job = TableAnalysisJob()
