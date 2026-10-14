"""
Retention Requirements task: find the retention period for each recommended archiving object in a retention
reference document (a records retention schedule) and say where in the document it was found.

  1. READ the document(s). Tables -- the usual shape of a retention schedule, in Word, Excel or CSV -- are read
     exactly, row by row: the columns are recognised by their headings (retention requirement, record series,
     code, trigger event, ...), so every entry knows its file, table/sheet and row. Documents without such a
     table (PDF, PowerPoint, prose) are read by the AI model, and each entry is then located in the text.
  2. MATCH. A schedule is organised by business records, not by SAP object, so for each archiving object the AI
     model is shown the schedule's entries (numbered) with the object's name and description, and picks the entry
     whose records the object holds -- or none. It can only pick an entry that exists, and the retention period
     is copied from that entry, never written by the model.
  3. WRITE results/Retention Requirements/Retention requirements.xlsx (saved automatically): per object the retention period, trigger, matched record
     series, confidence and a Reference note (file, location, series code), plus the entries that were read.

The job runs on background threads and talks to the UI through snapshot().
"""

import csv
import io
import logging
import re
import threading
import xml.etree.ElementTree as ET
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Literal, Optional

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from config import results_folder
import batching
import llm
import object_descriptions
import reference_doc
from completeness import Tally

logger = logging.getLogger(__name__)

OUTPUT_NAME = "Retention requirements.xlsx"
MAX_AI_CHUNKS = 12             # per unstructured document
MAX_ENTRIES = 3000

# ---------------------------------------------------------------------------
# Reading the documents
# ---------------------------------------------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _clean(value) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def _docx_tables(contents: bytes) -> list[dict]:
    """The tables of a Word file: [{"title": the paragraph just before it, "rows": [[cell text]]}]."""
    with zipfile.ZipFile(io.BytesIO(contents)) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    body = root.find(_W + "body")
    tables: list[dict] = []
    last_text = ""

    def para(p) -> str:
        return _clean("".join(t.text or "" for t in p.iter(_W + "t")))

    for child in (body if body is not None else []):
        if child.tag == _W + "p":
            text = para(child)
            if text:
                last_text = text
        elif child.tag == _W + "tbl":
            rows = []
            for tr in child.findall(_W + "tr"):
                rows.append([_clean(" ".join(para(p) for p in tc.iter(_W + "p"))) for tc in tr.findall(_W + "tc")])
            label = f"table {len(tables) + 1}"
            if len(last_text) <= 60 and not last_text.endswith((".", ":")):   # a heading, not a sentence
                label += f" ('{last_text}')"
            tables.append({"title": label, "rows": rows})
    return tables


def _xlsx_tables(contents: bytes) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    try:
        return [
            {"title": f"sheet '{ws.title}'",
             "rows": [[_clean(v) for v in row] for row in ws.iter_rows(values_only=True)]}
            for ws in wb.worksheets
        ]
    finally:
        wb.close()


def _csv_tables(contents: bytes) -> list[dict]:
    text = reference_doc._extract_plain_text(contents)
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [{"title": "", "rows": [[_clean(c) for c in row] for row in csv.reader(io.StringIO(text), dialect)]}]


def _tables_of(contents: bytes, filename: str) -> list[dict]:
    lower = filename.lower()
    if lower.endswith(".docx"):
        return _docx_tables(contents)
    if lower.endswith((".xlsx", ".xlsm")):
        return _xlsx_tables(contents)
    if lower.endswith(".csv"):
        return _csv_tables(contents)
    return []


# What each column of a schedule is, recognised by its heading. First match per field wins and a column is only
# used once; the record series name is the fallback for whatever is left.
_COLUMN_PATTERNS = [
    ("retention", r"retention|retain|keep|period|years"),
    ("trigger", r"trigger|event|starts?\b|begin|clock"),
    ("code", r"\bcode\b|series (no|#|id)|^id$|^no\.?$|ref(erence)? ?(no|#|id|number)"),
    ("description", r"descri|detail|scope|content|includes"),
    ("basis", r"citation|regulat|legal|statut|authority|basis|reference|\blaw\b|compliance|notes?\b|remarks?|comment"),
    ("format", r"format|medium|media"),
    ("function", r"function|group|department|business (area|unit)|\barea\b|domain|module|division"),
    ("name", r"series name|record series|record (type|class|category|name)|name|title|document type|category|record"),
]
_NOT_RETENTION = re.compile(r"trigger|event|start|format|basis|authority|citation|regulat|legal")


def _map_columns(header: list[str]) -> dict[str, int]:
    heads = [h.lower() for h in header]
    taken: set[int] = set()
    out: dict[str, int] = {}
    for field, pattern in _COLUMN_PATTERNS:
        rx = re.compile(pattern)
        options = [i for i, h in enumerate(heads) if i not in taken and h and rx.search(h)]
        if field == "retention":
            options = [i for i in options if not _NOT_RETENTION.search(heads[i])]
            # "Retention Requirement" / "Retention Period" beat a bare "Retention"
            options.sort(key=lambda i: 0 if re.search(r"requirement|period|years|time", heads[i]) else 1)
        if options:
            out[field] = options[0]
            taken.add(options[0])
    if "name" not in out:
        free = [i for i, h in enumerate(heads) if i not in taken and h]
        if free:
            out["name"] = free[0]
    return out


def _entries_from_table(table: dict, filename: str) -> list[dict]:
    """The schedule entries in one table, or [] if it isn't a retention schedule (no 'retention' heading)."""
    rows = table["rows"]
    header_at = next(
        (i for i, row in enumerate(rows[:25]) if sum(1 for c in row if c) >= 2 and any(re.search("retention", c, re.I) for c in row)),
        None,
    )
    if header_at is None:
        return []
    header = rows[header_at]
    cols = _map_columns(header)
    if "retention" not in cols or "name" not in cols:
        return []
    width = len(header)
    entries: list[dict] = []
    function = ""
    for offset, raw in enumerate(rows[header_at + 1:], start=header_at + 2):
        cells = [c for c in raw]
        if not any(cells):
            continue
        filled = [c for c in cells if c]
        if len(filled) == 1 and len(cells) <= width:     # a group heading row ("Administration")
            function = filled[0]
            continue
        if len(cells) > width:                            # merged cells give a row extra leading cells: align right
            cells = cells[-width:]
        cells = cells + [""] * (width - len(cells))
        if cells[cols["retention"]].lower() == header[cols["retention"]].lower():
            continue                                      # a repeated heading row (page break)

        def get(field: str) -> str:
            return cells[cols[field]] if field in cols else ""

        retention, name, code = get("retention"), get("name"), get("code")
        if not retention or not (name or code):
            continue
        where = f"{table['title']}, row {offset}" if table["title"] else f"row {offset}"
        entries.append({
            "file": filename, "location": where, "function": get("function") or function, "name": name,
            "code": code, "description": get("description"), "trigger": get("trigger"), "retention": retention,
            "basis": get("basis"), "format": get("format"), "method": "table",
        })
    return entries


# --- the AI path for documents that are not tables ---------------------------------------------------------

class _Entry(BaseModel):
    record_series: str = Field(description="Name of the record series / record category, as written.")
    code: str = Field(default="", description="Its code or number, if the document gives one.")
    description: str = Field(default="", description="What records it covers, briefly, as written.")
    trigger: str = Field(default="", description="The event the period starts from (e.g. 'end of fiscal year'), if stated.")
    retention: str = Field(description="The retention period exactly as written (e.g. '7 years', 'Permanent').")
    basis: str = Field(default="", description="The regulation or law the period comes from, if stated.")


class _Entries(BaseModel):
    entries: list[_Entry]


_EXTRACT_SYSTEM = (
    "You read records retention documents. List every record series / record category in the text together "
    "with its retention period. Copy the retention period, trigger event and regulation EXACTLY as written; never "
    "estimate or add one the text does not state. Skip anything without a retention period. If the text holds no "
    "retention rules, return an empty list."
)

_MARKER = re.compile(r"^=== (Page \d+|Slide \d+|Sheet: .+?) ===$")


def _locate(text: str, entry: dict) -> str:
    """Where in the text an AI-extracted entry sits: the nearest page / slide / sheet heading before the first line
    that mentions its record series (or code); '' if it cannot be found."""
    needles = [n.lower() for n in (entry["name"], entry["code"]) if n and len(n) >= 4]
    marker = ""
    for line in text.splitlines():
        m = _MARKER.match(line.strip())
        if m:
            marker = m.group(1).replace("Sheet: ", "sheet ")
            continue
        low = line.lower()
        if any(n in low for n in needles):
            return marker
    return ""


def _part_size(settings: llm.Settings) -> int:
    """Characters of a document one request reads: what the model takes in, and no more than its answer can carry
    (the same sizing the reference reviews use)."""
    return min(llm.budget(settings)["reference_chars"], max(2000, llm.max_output_tokens(settings) * 4))


def _entries_via_ai(text: str, filename: str, settings: llm.Settings, warnings: list[str], tally: Tally) -> list[dict]:
    """The document is read part by part (a part whose answer is cut off is split and asked again); what could not be
    read goes into *tally*. Raises the model's error if it is unusable (rejected key, rate limit)."""
    max_tokens = min(llm.max_output_tokens(settings), 4000)

    def ask(piece: str) -> "_Entries":
        return llm.structured(_EXTRACT_SYSTEM, "DOCUMENT TEXT:\n" + piece, _Entries, max_tokens=max_tokens, settings=settings)

    read = batching.run_text_parts(
        text, ask, size=_part_size(settings), max_parts=MAX_AI_CHUNKS, what=filename, tally=tally
    )
    if read.fatal is not None:
        raise read.fatal
    if not read.results and read.missing:
        warnings.append(f"{filename}: the model could not read it ({read.missing[0][1]}).")
        return []
    out: list[dict] = []
    seen: set[tuple] = set()
    for result in read.results:
        for e in result.entries:
            key = (e.record_series.strip().lower(), e.code.strip().lower(), e.retention.strip().lower())
            if not e.retention.strip() or not (e.record_series.strip() or e.code.strip()) or key in seen:
                continue
            seen.add(key)
            entry = {
                "file": filename, "location": "", "function": "", "name": _clean(e.record_series), "code": _clean(e.code),
                "description": _clean(e.description), "trigger": _clean(e.trigger), "retention": _clean(e.retention),
                "basis": _clean(e.basis), "format": "", "method": "AI",
            }
            entry["location"] = _locate(text, entry)
            out.append(entry)
    return out


def read_documents(files: list[tuple[bytes, str]], settings: llm.Settings, tally: Optional[Tally] = None, on_progress=None) -> dict:
    """Every retention entry in the uploaded documents: {"entries": [...], "documents": [{"name", "entries",
    "method"}], "warnings": [...]}. Unreadable documents are skipped with a warning; what the model's limits kept
    from being read is recorded in *tally*; raises llm.LLMError if the AI model is needed but cannot be used."""
    tally = tally if tally is not None else Tally()
    entries: list[dict] = []
    documents: list[dict] = []
    warnings: list[str] = []
    for contents, filename in files:
        if on_progress:
            on_progress(f"Reading {filename}…")
        found: list[dict] = []
        method = "table"
        try:
            for table in _tables_of(contents, filename):
                found.extend(_entries_from_table(table, filename))
            if not found:
                method = "AI"
                text = reference_doc._extract_text(contents, filename)
                if not text.strip():
                    warnings.append(f"{filename}: no text could be extracted (a scanned or image-only file?)")
                    continue
                if on_progress:
                    on_progress(f"{filename} has no retention table; the AI model is reading it…")
                found = _entries_via_ai(text, filename, settings, warnings, tally)
        except ValueError as exc:
            warnings.append(str(exc))
            continue
        except llm.LLMError:
            raise
        except Exception as exc:
            warnings.append(f"{filename}: could not be read ({exc})")
            continue
        documents.append({"name": filename, "entries": len(found), "method": method})
        if not found:
            warnings.append(f"{filename}: no retention entries were found in it.")
        entries.extend(found)
    for n, e in enumerate(entries[:MAX_ENTRIES], start=1):
        e["id"] = n
    return {"entries": entries[:MAX_ENTRIES], "documents": documents, "warnings": warnings}


# ---------------------------------------------------------------------------
# The retention period itself
# ---------------------------------------------------------------------------

_UNITS = {"year": 1.0, "yr": 1.0, "month": 1 / 12, "mo": 1 / 12, "week": 1 / 52, "day": 1 / 365}


def parse_retention(text: str) -> dict:
    """{"years": number or None, "kind": "Permanent" | "Fixed" | "After an event" | "See document"} from a period
    as the document words it ("7 Years", "Superseded + 6 Years", "Permanent"). The text itself is always kept
    as written; this only helps sort and compare."""
    t = (text or "").lower()
    if re.search(r"permanent|indefinite|forever", t):
        return {"years": None, "kind": "Permanent"}
    found = re.findall(r"(\d+(?:\.\d+)?)\s*(year|yr|month|mo|week|day)s?\b", t)
    event = bool(re.search(r"\+|after|superseded|life of|from|until|following|end of|termination", t))
    if len(found) == 1:
        n, unit = found[0]
        return {"years": round(float(n) * _UNITS[unit], 2), "kind": "After an event" if event else "Fixed"}
    return {"years": None, "kind": "See document"}


def reference_note(entry: dict) -> str:
    """Where the period was found, for the Reference column."""
    parts = [f"'{entry['file']}'"]
    if entry["location"]:
        parts.append(entry["location"])
    what = " ".join(p for p in (entry["code"], entry["name"]) if p)
    if entry["function"]:
        what += f" (function: {entry['function']})"
    if what:
        parts.append(what)
    note = " — ".join(parts)
    if entry["method"] == "AI":
        note += " [read by the AI model; check against the document]"
    return note


# ---------------------------------------------------------------------------
# Matching an archiving object to an entry
# ---------------------------------------------------------------------------

_STOP = {
    "and", "the", "for", "with", "from", "that", "this", "other", "all", "any", "are", "not", "data", "document",
    "documents", "record", "records", "related", "including", "information", "management", "general",
}


def usable_description(text: str) -> str:
    """An object description fit to show the model: an '(AI-suggested)' label is dropped, a '(not found)'
    placeholder becomes empty."""
    t = (text or "").strip()
    if t.lower().startswith("(ai-suggested)"):
        return t[len("(AI-suggested)"):].strip()
    return "" if t.startswith("(") else t


def _tokens(text: str) -> set[str]:
    words = re.findall(r"[a-z0-9]+", (text or "").lower())
    return {w[:-1] if len(w) > 4 and w.endswith("s") else w for w in words if len(w) >= 3 and w not in _STOP}


def compact(entry: dict) -> str:
    bits = [entry["code"], entry["function"], entry["name"], entry["description"][:110]]
    return " | ".join(b for b in bits if b)


class _Terms(BaseModel):
    terms: list[str] = Field(description="8-15 words or short phrases for the business records this archiving object holds.")


class _Choice(BaseModel):
    entry: int = Field(description="Number of the schedule entry whose records this archiving object holds, or 0 if none fits.")
    confidence: Literal["high", "medium", "low"] = Field(
        description="high: the entry clearly describes the same records; medium: related, probably right; low: a stretch."
    )
    reason: str = Field(description="One sentence: why that entry's records are what the archiving object holds.")
    also: list[int] = Field(default_factory=list, description="Up to 3 other entry numbers that could also apply, else empty.")


_TERMS_SYSTEM = (
    "You know SAP data archiving. For the given SAP archiving object, list 8-15 plain business words or short "
    "phrases for the kind of records it holds (for example 'invoices, journal entries, general ledger, accounts "
    "payable'), so they can be searched for in a records retention schedule."
)

_MATCH_SYSTEM = (
    "You match SAP archiving objects to the entries of a records retention schedule. An archiving object stores "
    "the data of one business area; the retention period of the schedule entry that covers those records is the "
    "retention period for the object. Choose the ONE numbered entry whose records are what the object holds. Use only "
    "the numbered entries given: never invent an entry or a period. A wrong match is worse than none, so if no "
    "entry clearly covers the object's data, answer entry 0. Do not match an object to an entry only because both "
    "are loosely 'business' or 'administrative', or because the data might be used in that area. Generic technical "
    "objects (links, logs, change documents, work items) belong only to an entry that covers system, audit-trail or "
    "workflow records; otherwise answer 0. Confidence: high only when the entry names the same kind of record the "
    "object stores; medium when the entry's area clearly covers the object's data without naming it; low when the "
    "connection is weak. Prefer the most specific entry. 'also' is normally empty: list another entry only if it genuinely "
    "applies too (for example another record type the object also stores), never as filler."
)


def _candidates(obj: str, description: str, entries: list[dict], budget: int, settings: llm.Settings,
                tally: Optional[Tally] = None) -> list[dict]:
    """The entries to show the model for one object: all of them if they fit the prompt budget, otherwise the best
    keyword matches (using the model's own list of what the object holds)."""
    total = sum(len(compact(e)) + 8 for e in entries)
    if total <= budget:
        return entries
    try:
        terms = llm.structured(_TERMS_SYSTEM, f"Archiving object: {obj}\nDescription: {description}", _Terms,
                               max_tokens=300, settings=settings).terms
    except llm.LLMError:
        terms = []
    wanted = _tokens(f"{obj.replace('_', ' ')} {description} {' '.join(terms)}")

    def score(e: dict) -> int:
        return 3 * len(wanted & _tokens(f"{e['function']} {e['name']}")) + len(wanted & _tokens(e["description"]))

    ranked = sorted(entries, key=lambda e: (-score(e), e["id"]))
    chosen, used = [], 0
    for e in ranked:
        size = len(compact(e)) + 8
        if used + size > budget and len(chosen) >= 12:
            break
        chosen.append(e)
        used += size
    if tally is not None:
        tally.note(
            f"The retention schedule has {len(entries)} entries, more than this model's prompt holds, so each archiving "
            f"object was matched against the {len(chosen)} or so most relevant ones (pre-selected by keywords). An object "
            "whose entry was not among them shows as no match; a model with a larger window sees them all."
        )
    return sorted(chosen, key=lambda e: e["id"])


def match_object(obj: str, description: str, entries: list[dict], settings: llm.Settings,
                 tally: Optional[Tally] = None) -> dict:
    """The schedule entry for one archiving object: {"entry": dict or None, "confidence", "reason", "also": [dicts]}."""
    candidates = _candidates(obj, description, entries, llm.budget(settings)["reference_chars"], settings, tally)
    listing = "\n".join(f"{n}. {compact(e)}" for n, e in enumerate(candidates, start=1))
    user = (f"Archiving object: {obj}\nDescription: {description or '(none given)'}\n\n"
            f"RETENTION SCHEDULE ENTRIES:\n{listing}")
    choice = llm.structured(_MATCH_SYSTEM, user, _Choice, max_tokens=500, settings=settings)

    def pick(n: int) -> Optional[dict]:
        return candidates[n - 1] if 1 <= n <= len(candidates) else None

    chosen = pick(choice.entry)
    also = [e for e in (pick(n) for n in choice.also[:3]) if e is not None and e is not chosen]
    return {"entry": chosen, "confidence": choice.confidence if chosen else "", "reason": choice.reason.strip(), "also": also}


# ---------------------------------------------------------------------------
# The job
# ---------------------------------------------------------------------------

class RetentionJob:
    def __init__(self):
        self._lock = threading.RLock()
        self._write_lock = threading.Lock()
        self._reset()

    def _reset(self) -> None:
        self._state = {
            "status": "idle",       # idle | running | done | error
            "total": 0,
            "completed": 0,
            "message": None,
            "documents": [],        # [{"name", "entries", "method"}]
            "output_path": None,
            "warnings": [],
        }
        self._records: dict[str, dict] = {}
        self._order: list[str] = []
        self._entries: list[dict] = []
        self._tally = Tally()
        self._cancel = False

    # -- public API -----------------------------------------------------------

    def snapshot(self) -> dict:
        with self._lock:
            state = dict(self._state)
            state["records"] = [dict(self._records[o]) for o in self._order]
            state["warnings"] = list(self._state["warnings"])
            state["documents"] = [dict(d) for d in self._state["documents"]]
            state["entry_count"] = len(self._entries)
        state["completeness"] = self._tally.summary()
        return state

    def entries(self) -> list[dict]:
        """What was read from the documents, for the 'check the extraction' preview."""
        with self._lock:
            return [dict(e) for e in self._entries]

    def is_active(self) -> bool:
        with self._lock:
            return self._state["status"] == "running"

    def start(self, objects: list[dict], documents: list[tuple[bytes, str]]) -> None:
        with self._lock:
            if self._state["status"] == "running":
                raise RuntimeError("A retention analysis is already running.")
            self._reset()
            self._state.update(status="running", total=len(objects), message="Starting…")
            for o in objects:
                self._order.append(o["object"])
                self._records[o["object"]] = self._blank(o)
        threading.Thread(target=self._run, args=(objects, documents), daemon=True, name="retention").start()

    def cancel(self) -> None:
        with self._lock:
            self._cancel = True

    # -- plumbing -------------------------------------------------------------

    @staticmethod
    def _blank(o: dict) -> dict:
        return {
            "object": o["object"], "description": o.get("description", ""), "state": "queued",
            "retention": "", "years": None, "period_type": "", "trigger": "", "series": "", "code": "", "function": "",
            "confidence": "", "reference": "", "basis": "", "reason": "", "alternatives": [], "note": "Waiting.",
        }

    def _set(self, **changes) -> None:
        with self._lock:
            self._state.update(changes)

    def _warn(self, text: str) -> None:
        with self._lock:
            if text not in self._state["warnings"]:
                self._state["warnings"].append(text)

    def _fail(self, message: str) -> None:
        with self._lock:
            self._state.update(status="error", message=message)

    def _run(self, objects: list[dict], documents: list[tuple[bytes, str]]) -> None:
        try:
            settings = llm.load_settings()
            problem = llm.not_configured_message(settings)
            if problem:
                self._fail(problem)
                return
            read = read_documents(documents, settings, self._tally, on_progress=lambda m: self._set(message=m))
            with self._lock:
                self._entries = read["entries"]
                self._state["documents"] = read["documents"]
                self._state["warnings"].extend(read["warnings"])
            if not read["entries"]:
                self._fail("No retention entries were found in the document(s). " + " ".join(read["warnings"]))
                return

            self._set(message="Looking up the descriptions of the archiving objects…")
            known = {o["object"]: o["description"] for o in objects if o.get("description")}
            try:
                described = object_descriptions.resolve([o["object"] for o in objects], known=known,
                                                        settings=settings, allow_sap=False)
            except Exception:
                logger.exception("Could not resolve object descriptions")
                described = known
            with self._lock:
                for o in objects:
                    self._records[o["object"]]["description"] = usable_description(described.get(o["object"], ""))                         or self._records[o["object"]]["description"]

            self._set(message="Matching the archiving objects to the retention entries…")
            workers = max(1, llm.max_workers(settings))
            with ThreadPoolExecutor(max_workers=workers) as pool:
                futures = [pool.submit(self._one, o["object"], read["entries"], settings) for o in objects]
                for f in futures:
                    f.result()
        except llm.LLMError as exc:
            logger.warning("Retention analysis stopped: %s", exc)
            if self._entries:
                self._save()   # keep what was matched so far; saved before the run counts as stopped
            self._fail(str(exc))
            return
        except Exception as exc:
            logger.exception("Retention analysis failed")
            if self._entries:
                self._save()
            self._fail(f"Unexpected error: {exc}")
            return
        self._save()   # the file is complete before the run counts as done
        with self._lock:
            self._state.update(status="done", message="Cancelled." if self._cancel else "Finished.")

    def _one(self, obj: str, entries: list[dict], settings: llm.Settings) -> None:
        with self._lock:
            if self._cancel:
                rec = self._records[obj]
                if rec["state"] == "queued":
                    rec.update(state="skipped", note="Cancelled before it ran.")
                self._state["completed"] = self._count()
                return
            rec = self._records[obj]
            rec.update(state="running", note="Matching…")
            description = rec["description"]
        try:
            result = match_object(obj, description, entries, settings, self._tally)
        except (llm.LLMNotConfiguredError, llm.LLMAuthError):
            raise
        except llm.LLMRateLimitError as exc:
            if exc.daily:
                raise
            self._finish_failed(obj, str(exc))
            return
        except llm.LLMError as exc:
            self._finish_failed(obj, str(exc))
            return
        except Exception as exc:
            logger.exception("Retention matching failed for %s", obj)
            self._finish_failed(obj, f"Unexpected error: {exc}")
            return

        entry = result["entry"]
        with self._lock:
            rec = self._records[obj]
            if entry is None:
                rec.update(state="no_match", reason=result["reason"],
                           note="No entry in the retention document covers this archiving object.")
            else:
                parsed = parse_retention(entry["retention"])
                rec.update(
                    state="done", retention=entry["retention"], years=parsed["years"], period_type=parsed["kind"],
                    trigger=entry["trigger"], series=entry["name"], code=entry["code"], function=entry["function"],
                    confidence=result["confidence"], reference=reference_note(entry), basis=entry["basis"],
                    reason=result["reason"], note="",
                    alternatives=[{
                        "series": f"{e['code']} {e['name']}".strip(), "retention": e["retention"], "trigger": e["trigger"],
                        "reference": reference_note(e),
                    } for e in result["also"]],
                )
            self._state["completed"] = self._count()
        self._save()

    def _finish_failed(self, obj: str, message: str) -> None:
        with self._lock:
            self._records[obj].update(state="failed", note=message)
            self._state["completed"] = self._count()
        self._save()

    def _count(self) -> int:
        return sum(1 for r in self._records.values() if r["state"] in ("done", "no_match", "failed", "skipped"))

    # -- output ---------------------------------------------------------------

    def _save(self) -> None:
        try:
            self._write_workbook()
        except Exception as exc:
            logger.exception("Could not write the Retention requirements workbook")
            self._warn(f"Could not save the Excel file: {exc}")

    def _write_workbook(self) -> None:
        with self._write_lock:
            with self._lock:
                records = [dict(self._records[o]) for o in self._order]
                entries = [dict(e) for e in self._entries]
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Retention requirements"
            ws.append([
                "Archiving Object", "Description", "Retention Period", "Retention (years)", "Period Type", "Trigger Event",
                "Matched Record Series", "Series Code", "Function", "Confidence", "Reference (where the period was found)",
                "Basis / Regulation", "Why this match", "Other possible matches", "Note",
            ])
            for r in records:
                others = "; ".join(f"{a['series']}: {a['retention']} ({a['reference']})" for a in r["alternatives"])
                ws.append([
                    r["object"], r["description"], r["retention"], r["years"] if r["years"] is not None else "",
                    r["period_type"], r["trigger"], r["series"], r["code"], r["function"], r["confidence"].capitalize(),
                    r["reference"], r["basis"], r["reason"], others,
                    r["note"] if r["state"] != "done" else "",
                ])
            self._style(ws)
            ws.freeze_panes = "B2"

            read = wb.create_sheet("Entries read")
            read.append(["#", "File", "Location", "Function", "Record Series", "Series Code", "Description", "Trigger Event",
                         "Retention", "Basis / Regulation", "Read by"])
            for e in entries:
                read.append([e["id"], e["file"], e["location"], e["function"], e["name"], e["code"], e["description"],
                             e["trigger"], e["retention"], e["basis"], e["method"]])
            self._style(read)
            notes = self._tally.sentences(limit=30)
            if notes:
                sheet = wb.create_sheet("Notes")
                sheet.append(["What the model's limits kept this result from covering"])
                for line in notes:
                    sheet.append([line])
                self._style(sheet)
                sheet.column_dimensions["A"].width = 120
            path = self._save_workbook(wb)
        self._set(output_path=str(path))

    @staticmethod
    def _style(ws) -> None:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        widths: dict[int, int] = {}
        for row in ws.iter_rows(values_only=True):
            for i, v in enumerate(row, start=1):
                widths[i] = max(widths.get(i, 8), min(len(str(v)) + 2 if v is not None else 0, 60))
        for i, w in widths.items():
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    @staticmethod
    def _save_workbook(wb) -> Path:
        folder = results_folder("RETENTION")
        stem, suffix = Path(OUTPUT_NAME).stem, Path(OUTPUT_NAME).suffix
        for path in [folder / OUTPUT_NAME] + [folder / f"{stem} ({n}){suffix}" for n in range(2, 10)]:
            try:
                wb.save(path)
                return path
            except PermissionError:   # open in Excel
                continue
        raise PermissionError("Could not write the Retention requirements workbook (is it open in Excel?).")


job = RetentionJob()
