"""
Archiving object analysis: for each recommended archiving object, find

  * its ARCHIVING CONDITIONS, from three sources, each labelled in the output:
      1. SARA's information (`i`) button -- the object's SAP Help Portal page (and its "Checks"-type
         sub-pages), found by pressing the button in SARA and reading the page SAP opens;
      2. the DVM Guide;
      3. SAP for Me (SAP Notes / Knowledge Base Articles; SAP Community as a fallback) -- SLOW, so it is a
         separate second step the user starts after previewing 1 and 2 (start_sap_for_me()); it runs in
         the background, needs only its own browser and the AI model (not the SAP GUI), and adds its
         conditions to the same results and workbook as it finds them;
  * the archiving objects that must be ARCHIVED BEFORE it, from SARA's network (table ARCH_NET -- the data
    behind the "Network Graphic" button), followed through all levels and put in archiving order.

Everything is written to  output/Archiving object analysis.xlsx : a Summary sheet and one sheet per object,
with a clearly marked "Archive these objects BEFORE ..." block and the conditions table. The workbook is
rewritten after every object, so finished objects survive a failure or cancel.

The jobs run on background threads; the UI polls snapshot() and reads details() for the preview.
"""

import io
import logging
import re
import threading
import time
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

import dvm_guide
import help_browser
import helpportal
import llm
import object_descriptions
import sap_for_me
import transactions.db02 as db02
import transactions.sara_info as sara_info
from config import OUTPUT_DIR

logger = logging.getLogger(__name__)

OUTPUT_NAME = "Archiving object analysis.xlsx"
_OBJECT_NAME = re.compile(r"^[A-Z0-9_/]{2,30}$")

SRC_SARA = "SARA information (SAP Help Portal)"
SRC_DVM = "DVM Guide"
SRC_SFM = "SAP for Me"


# ---------------------------------------------------------------------------
# Input: the recommended archiving objects
# ---------------------------------------------------------------------------

def _norm(text) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def parse_objects(contents: bytes) -> list[dict]:
    """Distinct archiving objects, in order, from a recommended-objects workbook: [{"object", "description"}].

    Uses the last sheet whose title contains "recommended" and that has an "Archiving Object" column (the
    scored workbook's "Recommended" sheet, or "Recommended (with Reference Doc)" after the reference review);
    otherwise the last sheet with that column. Blank cells, placeholders like "(no archiving object found)"
    and invalid names are skipped. Raises ValueError with a readable message."""
    workbook = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    try:
        candidates = []
        for ws in workbook.worksheets:
            header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if not header:
                continue
            cols = {_norm(c): i for i, c in enumerate(header) if c is not None}
            if "archivingobject" in cols:
                candidates.append((ws, cols))
        if not candidates:
            raise ValueError(
                "No 'Archiving Object' column found. Use the workbook saved by Find Archiving Objects "
                "(archiving_objects_scored.xlsx or archiving_objects_with_reference.xlsx)."
            )
        preferred = [c for c in candidates if "recommended" in c[0].title.lower()]
        ws, cols = (preferred or candidates)[-1]
        obj_idx, desc_idx = cols["archivingobject"], cols.get("objectdescription")

        out: list[dict] = []
        seen: set[str] = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            value = str(row[obj_idx] if obj_idx < len(row) and row[obj_idx] is not None else "").strip().upper()
            if not value or value.startswith("(") or not _OBJECT_NAME.match(value) or value in seen:
                continue
            seen.add(value)
            desc = ""
            if desc_idx is not None and desc_idx < len(row) and row[desc_idx] is not None:
                desc = str(row[desc_idx]).strip()
            out.append({"object": value, "description": desc})
        if not out:
            raise ValueError("The 'Archiving Object' column has no archiving objects.")
        return out
    finally:
        workbook.close()


# ---------------------------------------------------------------------------
# Dependencies: which objects must be archived first (ARCH_NET)
# ---------------------------------------------------------------------------

def build_network(pairs: list[dict]) -> dict[str, set[str]]:
    """{OBJECT: {objects that must be archived before it}} from ARCH_NET rows."""
    net: dict[str, set[str]] = {}
    for p in pairs:
        net.setdefault(p["object"], set())
        if p["previous"]:
            net[p["object"]].add(p["previous"])
    return net


def prerequisites(obj: str, net: dict[str, set[str]]) -> list[dict]:
    """Every object that must be archived before *obj*, direct or indirect, in ARCHIVING ORDER.

    Each item: {"object", "step", "direct", "required_by"}: step 1 is archived first; "direct" says whether
    *obj* itself lists it as its predecessor; "required_by" names the objects (within the chain, or *obj*)
    that list it directly, i.e. why it is needed. Cycles in the data cannot loop: every object is visited once."""
    ancestors: set[str] = set()
    stack = list(net.get(obj, ()))
    while stack:
        node = stack.pop()
        if node in ancestors or node == obj:
            continue
        ancestors.add(node)
        stack.extend(net.get(node, ()))
    if not ancestors:
        return []

    within = ancestors | {obj}
    needs = {n: {p for p in net.get(n, ()) if p in ancestors} for n in ancestors}   # n needs p first
    required_by: dict[str, list[str]] = {a: [] for a in ancestors}
    for n in within:
        for p in net.get(n, ()):
            if p in ancestors:
                required_by[p].append(n)

    # topological layering: step 1 = nothing in the chain has to come before it
    steps: dict[str, int] = {}
    remaining = set(ancestors)
    step = 0
    while remaining:
        step += 1
        ready = sorted(n for n in remaining if not (needs[n] & remaining))
        if not ready:           # a cycle: release the rest together rather than loop forever
            ready = sorted(remaining)
        for n in ready:
            steps[n] = step
        remaining -= set(ready)

    direct = net.get(obj, set())
    return [
        {"object": n, "step": steps[n], "direct": n in direct, "required_by": sorted(required_by[n])}
        for n in sorted(ancestors, key=lambda n: (steps[n], n))
    ]


# ---------------------------------------------------------------------------
# Conditions: extracting them from each source's text
# ---------------------------------------------------------------------------

class _Condition(BaseModel):
    condition: str = Field(description="One archiving condition as a short, self-contained sentence.")
    detail: str = Field(default="", description="A key setting, check, status or number it refers to (short), or empty.")
    block: int = Field(default=0, description="Number of the text block it comes from, or 0 if unclear.")


class _Conditions(BaseModel):
    conditions: list[_Condition] = Field(default_factory=list)


_EXTRACT_SYSTEM = (
    "You extract the ARCHIVING CONDITIONS of one SAP archiving object from documentation excerpts. An archiving "
    "condition is something that must be true, or be done, before data can be archived or deleted with that "
    "object: checks the write/delete program performs, required document or processing status, minimum "
    "residence or retention times, required Customizing, other documents or archiving objects that must be "
    "archived or processed first, and business or technical blockers. Use ONLY the text given - never add "
    "conditions from your own knowledge. Skip general descriptions of what the object is or how it works. "
    "Write each condition as one short, self-contained sentence, merge duplicates, and give the number of the "
    "text block it came from. If the text contains no archiving conditions, return an empty list."
)


def extract_conditions(settings: llm.Settings, obj: str, description: str, blocks: list[dict]) -> list[dict]:
    """[{"condition", "detail", "source"}] from *blocks* ({"label", "text"}); each condition's source is the
    label of the block the model says it came from. Raises llm.LLMError."""
    if not blocks:
        return []
    parts = [f"Archiving object: {obj}" + (f" ({description})" if description else ""), "", "TEXT BLOCKS:"]
    for n, b in enumerate(blocks, start=1):
        parts.append(f"[{n}] {b['label']}\n---\n{b['text']}\n---")
    result = llm.structured(_EXTRACT_SYSTEM, "\n".join(parts), _Conditions, max_tokens=1500, settings=settings)

    all_labels = "; ".join(dict.fromkeys(b["label"] for b in blocks))
    out: list[dict] = []
    seen: set[str] = set()
    for c in result.conditions:
        key = _norm(c.condition)
        if not key or key in seen:
            continue
        seen.add(key)
        label = blocks[c.block - 1]["label"] if 1 <= c.block <= len(blocks) else all_labels
        out.append({"condition": c.condition.strip(), "detail": c.detail.strip(), "source": label})
    return out


# ---------------------------------------------------------------------------
# The job
# ---------------------------------------------------------------------------

class ObjectAnalysisJob:
    def __init__(self):
        self._lock = threading.RLock()
        self._write_lock = threading.Lock()
        self._sfm_cancel = False
        self._reset()

    def _reset(self) -> None:
        self._state = {
            "status": "idle",     # idle | running | done | error
            "total": 0,
            "completed": 0,
            "message": None,
            "current": None,      # {"object", "step"}
            "records": [],        # one summary per object finished
            "output_path": None,
            "warnings": [],
        }
        self._results: list[dict] = []
        self._cancel = False
        self._last_help_url: Optional[str] = None
        # the SAP for Me step (second, optional, runs in the background after the first step)
        self._sfm_state = {"status": "idle", "total": 0, "completed": 0, "message": None, "current": None}

    # -- public ---------------------------------------------------------------

    def snapshot(self) -> dict:
        with self._lock:
            state = dict(self._state)
            state["records"] = [dict(r) for r in self._state["records"]]
            state["warnings"] = list(self._state["warnings"])
            state["sap_for_me"] = dict(self._sfm_state)
            return state

    def details(self) -> list[dict]:
        """The full results for the on-screen preview: conditions (with sources), objects to archive first,
        what each source returned."""
        with self._lock:
            out = []
            for r in sorted(self._results, key=lambda r: r["index"]):
                out.append({
                    "object": r["object"],
                    "description": r.get("description", ""),
                    "error": r.get("error", ""),
                    "sara_url": r.get("sara_url", ""),
                    "network_known": r.get("network_known", False),
                    "sources": dict(r["sources"]),
                    "conditions": [dict(c) for c in r["conditions"]],
                    "prerequisites": [dict(p) for p in r["prerequisites"]],
                })
            return out

    def is_active(self) -> bool:
        """True while the first step or the SAP for Me step is running (a new run would reset the results)."""
        with self._lock:
            return self._state["status"] == "running" or self._sfm_state["status"] == "running"

    def start(self, objects: list[dict]) -> None:
        with self._lock:
            if self.is_active():
                raise RuntimeError("An archiving object analysis is already running.")
            self._reset()
            self._state.update(status="running", total=len(objects), message="Starting…")
        threading.Thread(target=self._run, args=(objects,), daemon=True, name="object-analysis").start()

    def cancel(self) -> None:
        with self._lock:
            self._cancel = True

    def start_sap_for_me(self) -> None:
        """Add SAP for Me conditions to the finished results, in the background (own browser + AI model only)."""
        with self._lock:
            if self._state["status"] != "done" or not self._results:
                raise RuntimeError("Run the analysis first; SAP for Me is checked afterwards.")
            if self._sfm_state["status"] == "running":
                raise RuntimeError("SAP for Me is already being checked.")
            targets = [r["object"] for r in sorted(self._results, key=lambda r: r["index"])]
            self._sfm_cancel = False
            self._sfm_state = {
                "status": "running", "total": len(targets), "completed": 0, "message": "Starting…", "current": None,
            }
        threading.Thread(target=self._run_sap_for_me, args=(targets,), daemon=True, name="object-analysis-sfm").start()

    def cancel_sap_for_me(self) -> None:
        with self._lock:
            self._sfm_cancel = True

    # -- plumbing -------------------------------------------------------------

    def _set(self, **changes) -> None:
        with self._lock:
            self._state.update(changes)

    def _step(self, obj: str, step: str) -> None:
        self._set(current={"object": obj, "step": step}, message=f"{obj}: {step}")

    def _cancelled(self) -> bool:
        with self._lock:
            return self._cancel

    def _warn(self, text: str) -> None:
        with self._lock:
            if text not in self._state["warnings"]:
                self._state["warnings"].append(text)

    # -- the run --------------------------------------------------------------

    def _run(self, objects: list[dict]) -> None:
        reader = helpportal.HelpPortalReader()
        self._last_help_url = None
        try:
            settings = llm.load_settings()
            problem = llm.not_configured_message(settings)
            if problem:
                raise RuntimeError(f"The AI model is not available: {problem}")

            self._set(message="Reading the archiving-object network (table ARCH_NET)…")
            network = db02.run_get_archiving_network()
            net: Optional[dict] = None
            if network["status"] == "ok":
                net = build_network(network["pairs"])
            else:
                self._warn(f"Could not read the archiving-object network: {network.get('message')}. "
                           "Objects to archive first are not shown.")

            known = {o["object"]: o["description"] for o in objects if o["description"]}
            for number, item in enumerate(objects, start=1):
                if self._cancelled():
                    break
                obj = item["object"]
                try:
                    result = self._analyse(obj, known.get(obj, ""), net, settings, reader, known)
                except Exception as exc:
                    logger.exception("Archiving object analysis failed for %s", obj)
                    result = {"object": obj, "description": known.get(obj, ""), "error": str(exc),
                              "conditions": [], "prerequisites": [], "sources": {}}
                self._record(result, number)

            with self._lock:
                self._state["status"] = "done"
                self._state["current"] = None
                self._state["message"] = "Cancelled." if self._cancel else "Finished."
        except Exception as exc:
            logger.exception("Archiving object analysis failed")
            with self._lock:
                self._state["status"] = "error"
                self._state["message"] = str(exc)
        finally:
            reader.close()
            self._close_help_tab()

    @staticmethod
    def _record_for(result: dict) -> dict:
        by_source = {k: 0 for k in (SRC_SARA, SRC_DVM, SRC_SFM)}
        for c in result["conditions"]:
            for k in by_source:
                if c["source"].startswith(k):
                    by_source[k] += 1
        # "failed" when nothing at all could be read (an object SARA rejects, every source unavailable)
        read_any = any(
            str(result["sources"].get(src, "")).startswith("read") for src in (SRC_SARA, SRC_DVM, SRC_SFM)
        )
        return {
            "object": result["object"],
            "description": result.get("description", ""),
            "state": "failed" if (result.get("error") or not read_any) else "done",
            "conditions": len(result["conditions"]),
            "by_source": by_source,
            "archive_first": [p["object"] for p in result["prerequisites"]],
            "note": result.get("error") or "; ".join(
                f"{name}: {info}"
                for name, info in result["sources"].items()
                if info and not info.startswith(("read", "not checked yet"))
            ),
        }

    def _record(self, result: dict, number: int) -> None:
        with self._lock:
            self._results.append({**result, "index": number})
            self._state["records"].append(self._record_for(result))
            self._state["completed"] = len(self._state["records"])
        self._save_workbook()

    def _save_workbook(self) -> None:
        try:
            self._write_workbook()
        except Exception as exc:
            logger.exception("Could not write the analysis workbook")
            self._warn(f"Could not save the Excel file: {exc}")

    # -- one object -----------------------------------------------------------

    def _close_help_tab(self) -> None:
        url = self._last_help_url
        if url:
            if help_browser.close_help_page(url):
                self._last_help_url = None

    def _analyse(self, obj, description, net, settings, reader, known) -> dict:
        budget = llm.budget(settings)
        conditions: list[dict] = []
        sources: dict[str, str] = {}

        # 1. SARA information button -> the SAP Help Portal page
        self._step(obj, "SARA: opening the information page")
        self._close_help_tab()                       # the previous object's page is closed first
        blocks, sara_note, sara_url = [], "", ""
        before = help_browser.help_urls_open()
        pressed = sara_info.open_information_page(obj)
        if pressed["status"] != "ok":
            sara_note = pressed["message"]
        else:
            description = description or pressed["description"]
            url = help_browser.wait_for_new_help_page(before)
            if not url:
                sara_note = ("SARA's information button did not open a SAP Help Portal page in the browser "
                             "(or its address could not be read).")
            else:
                self._last_help_url = url
                self._close_help_tab()               # close it straight away; the page is read headlessly
                sara_url = url
                self._step(obj, "Reading the Help Portal page")
                page = reader.read(url, max_chars=budget["article_chars"] * 2)
                if page["problem"]:
                    sara_note = page["problem"]
                else:
                    blocks = [{"label": f"{SRC_SARA} - {p['title']}", "text": p["text"]} for p in page["pages"]]
                    sources[SRC_SARA] = f"read {len(page['pages'])} page(s): " + " | ".join(p["title"] for p in page["pages"])
        if sara_note:
            sources[SRC_SARA] = sara_note
        conditions += self._conditions(settings, obj, description, blocks, sources, SRC_SARA)

        # 2. DVM Guide
        if self._cancelled():
            sources[SRC_SFM] = "not checked yet"
            return self._result(obj, description, conditions, net, known, sources, sara_url)
        self._step(obj, "Searching the DVM Guide")
        hits = dvm_guide.sections_naming(obj, limit=4, window=max(300, budget["dvm_excerpt_chars"] // 3))
        dvm_blocks = [
            {"label": f"{SRC_DVM} - section on {', '.join(h['tables'][:5])}", "text": re.sub(r"\s+", " ", h["excerpt"])}
            for h in hits
        ]
        if not dvm_blocks:
            sources[SRC_DVM] = "the DVM Guide does not mention this object" if dvm_guide.load_index() else "the DVM Guide is not available"
        else:
            sources[SRC_DVM] = f"read {len(dvm_blocks)} section(s)"
        conditions += self._conditions(settings, obj, description, dvm_blocks, sources, SRC_DVM)

        # SAP for Me is slow: it is a separate step the user starts after previewing these results.
        sources[SRC_SFM] = "not checked yet"
        return self._result(obj, description, conditions, net, known, sources, sara_url)

    def _conditions(self, settings, obj, description, blocks, sources, label) -> list[dict]:
        if not blocks:
            return []
        self._step(obj, f"Extracting the conditions ({label})")
        return self._extract_into(settings, obj, description, blocks, sources, label)

    def _extract_into(self, settings, obj, description, blocks, sources, label) -> list[dict]:
        try:
            found = extract_conditions(settings, obj, description, blocks)
        except llm.LLMError as exc:
            sources[label] = f"{sources.get(label, '')} - could not extract conditions: {exc}".strip(" -")
            return []
        sources[label] = f"{sources.get(label, '')} - {len(found)} condition(s)".strip(" -")
        return found

    def _sap_for_me_blocks(self, obj, settings, sfm, sources) -> list[dict]:
        if not sfm["tried"]:
            sfm["tried"] = True
            try:
                sfm["session"] = sap_for_me.open_session()
                if sfm["session"] is None:
                    sfm["note"] = "no SAP for Me credentials are saved"
            except Exception as exc:
                sfm["note"] = f"sign-in failed: {sap_for_me._describe(exc)}"
        if sfm["session"] is None:
            sources[SRC_SFM] = f"not used: {sfm['note']}"
            return []
        try:
            articles, problem = sfm["session"].fetch_articles(
                f"{obj} archiving conditions prerequisites", obj, settings, subject=obj
            )
        except Exception as exc:
            sources[SRC_SFM] = f"search failed: {sap_for_me._describe(exc)}"
            return []
        if not articles:
            sources[SRC_SFM] = f"nothing found: {problem}"
            return []
        sources[SRC_SFM] = f"read {len(articles)} article(s)"
        return [{"label": f"{SRC_SFM} - {kind}: {title}", "text": text} for title, kind, text in articles]

    def _result(self, obj, description, conditions, net, known, sources, sara_url) -> dict:
        self._step(obj, "Following the dependencies")
        if net is None:
            prereqs = []
            sources["Dependencies (ARCH_NET)"] = "not available"
        else:
            prereqs = prerequisites(obj, net)
            sources["Dependencies (ARCH_NET)"] = (
                f"read {len(prereqs)} object(s) to archive first" if prereqs else "read none to archive first"
            )
        names = [p["object"] for p in prereqs]
        descriptions: dict[str, str] = {}
        if names:
            try:
                descriptions = object_descriptions.resolve(names, known=known)
            except Exception as exc:
                logger.warning("Could not resolve descriptions: %s", exc)
        for p in prereqs:
            p["description"] = descriptions.get(p["object"], "")
        return {
            "object": obj, "description": description, "conditions": conditions,
            "prerequisites": prereqs, "sources": sources, "sara_url": sara_url,
            "network_known": net is not None,
        }

    # -- SAP for Me (second, background step) -------------------------------------

    def _set_sfm(self, **changes) -> None:
        with self._lock:
            self._sfm_state.update(changes)

    def _run_sap_for_me(self, targets: list[str]) -> None:
        sfm: dict = {"session": None, "tried": False, "note": ""}
        try:
            settings = llm.load_settings()
            problem = llm.not_configured_message(settings)
            if problem:
                raise RuntimeError(f"The AI model is not available: {problem}")
            for number, obj in enumerate(targets, start=1):
                with self._lock:
                    if self._sfm_cancel:
                        break
                self._set_sfm(message=f"{obj}: searching SAP for Me", current=obj)
                with self._lock:
                    result = next((r for r in self._results if r["object"] == obj), None)
                    description = result.get("description", "") if result else ""
                if result is None:
                    continue
                local: dict[str, str] = {}
                try:
                    blocks = self._sap_for_me_blocks(obj, settings, sfm, local)
                    self._set_sfm(message=f"{obj}: extracting the conditions")
                    found = self._extract_into(settings, obj, description, blocks, local, SRC_SFM) if blocks else []
                except Exception as exc:
                    logger.exception("SAP for Me step failed for %s", obj)
                    found, local = [], {SRC_SFM: f"failed: {exc}"}
                with self._lock:
                    result["conditions"] = [c for c in result["conditions"] if not c["source"].startswith(SRC_SFM)] + found
                    result["sources"][SRC_SFM] = local.get(SRC_SFM, "nothing found")
                    refreshed = self._record_for(result)
                    for i, rec in enumerate(self._state["records"]):
                        if rec["object"] == obj:
                            self._state["records"][i] = refreshed
                    self._sfm_state["completed"] = number
                self._save_workbook()
            self._close_sfm_session(sfm)      # close the browser BEFORE reporting done
            with self._lock:
                self._sfm_state.update(
                    status="done", current=None,
                    message="Cancelled." if self._sfm_cancel else "SAP for Me finished.",
                )
        except Exception as exc:
            logger.exception("SAP for Me step failed")
            self._close_sfm_session(sfm)
            with self._lock:
                self._sfm_state.update(status="error", message=str(exc), current=None)
        finally:
            self._close_sfm_session(sfm)

    @staticmethod
    def _close_sfm_session(sfm: dict) -> None:
        session, sfm["session"] = sfm.get("session"), None
        if session is not None:
            try:
                session.close()
            except Exception:
                pass

    # -- output ---------------------------------------------------------------

    _ORANGE = PatternFill("solid", fgColor="F59E0B")
    _BLUE = PatternFill("solid", fgColor="1D4ED8")
    _GREY = PatternFill("solid", fgColor="E5E7EB")
    _AMBER_LIGHT = PatternFill("solid", fgColor="FEF3C7")

    def _write_workbook(self) -> None:
        with self._write_lock:           # the two steps may both finish an object at the same moment
            self._write_workbook_locked()

    def _write_workbook_locked(self) -> None:
        wb = openpyxl.Workbook()
        summary = wb.active
        summary.title = "Summary"
        headers = ["Archiving object", "Description", "ARCHIVE BEFORE (in order)", "Directly requires",
                   "Conditions", SRC_SARA, SRC_DVM, SRC_SFM, "Result / notes"]
        summary.append(headers)
        for cell in summary[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        summary["C1"].fill = self._ORANGE

        with self._lock:
            results = [
                {**r, "conditions": list(r["conditions"]), "sources": dict(r["sources"])}
                for r in sorted(self._results, key=lambda r: r["index"])
            ]
            records = {r["object"]: dict(r) for r in self._state["records"]}
        for res in results:
            rec = records.get(res["object"], {})
            prereqs = res["prerequisites"]
            order = ", ".join(p["object"] for p in prereqs) if prereqs else (
                "None" if res.get("network_known") else "(network not available)"
            )
            direct = ", ".join(p["object"] for p in prereqs if p["direct"]) or ("None" if res.get("network_known") else "")
            by = rec.get("by_source", {})
            summary.append([
                res["object"], res.get("description", ""), order, direct, len(res["conditions"]),
                by.get(SRC_SARA, 0), by.get(SRC_DVM, 0), by.get(SRC_SFM, 0),
                rec.get("note") or "OK",
            ])
            summary.cell(row=summary.max_row, column=3).fill = self._AMBER_LIGHT
        for row in summary.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col, width in zip("ABCDEFGHI", (22, 38, 52, 34, 12, 20, 14, 14, 60)):
            summary.column_dimensions[col].width = width
        summary.freeze_panes = "A2"

        used = {"Summary"}
        for res in results:
            self._object_sheet(wb.create_sheet(self._sheet_name(res["object"], used)), res)
        self._set(output_path=str(self._save(wb)))

    def _object_sheet(self, ws, res: dict) -> None:
        wrap = Alignment(vertical="top", wrap_text=True)
        for col, width in zip("ABCDE", (11, 62, 46, 40, 34)):
            ws.column_dimensions[col].width = width

        ws.append([f"Archiving object: {res['object']}"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([res.get("description", "")])
        ws.append([])
        ws.append(["Sources checked"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for name, info in res["sources"].items():
            ws.append(["", f"{name}: {info}"])
            ws.cell(row=ws.max_row, column=2).alignment = wrap
        if res.get("sara_url"):
            ws.append(["", f"SAP Help Portal page: {res['sara_url']}"])
        if res.get("error"):
            ws.append(["", f"ERROR: {res['error']}"])
            ws.cell(row=ws.max_row, column=2).font = Font(bold=True, color="B91C1C")
        ws.append([])

        # --- dependencies: the call-out
        band = f"ARCHIVE THESE OBJECTS BEFORE {res['object']}"
        ws.append([band])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws.cell(row=r, column=1).fill = self._ORANGE
        ws.cell(row=r, column=1).font = Font(bold=True, size=12, color="FFFFFF")
        prereqs = res["prerequisites"]
        if not res.get("network_known"):
            ws.append(["", "The archiving-object network (table ARCH_NET) could not be read, so this is unknown."])
        elif not prereqs:
            ws.append(["", f"None - no archiving object has to be archived before {res['object']} (per the SARA network, table ARCH_NET)."])
            ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
        else:
            ws.append(["Step", "Archiving object", "Description", "Relationship", "Required before"])
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
                cell.fill = self._GREY
            for p in prereqs:
                targets = ", ".join(p["required_by"]) if p["required_by"] else ""
                ws.append([
                    p["step"], p["object"], p.get("description", ""),
                    f"Direct - {res['object']} requires it first" if p["direct"] else "Indirect - needed further up the chain",
                    targets,
                ])
                for cell in ws[ws.max_row]:
                    cell.alignment = wrap
                    cell.fill = self._AMBER_LIGHT
                ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
            ws.append(["", f"Archive in step order (step 1 first); {res['object']} itself comes after the last step."])
            ws.cell(row=ws.max_row, column=2).font = Font(italic=True)
        ws.append([])

        # --- conditions
        ws.append(["ARCHIVING CONDITIONS"])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws.cell(row=r, column=1).fill = self._BLUE
        ws.cell(row=r, column=1).font = Font(bold=True, size=12, color="FFFFFF")
        if not res["conditions"]:
            ws.append(["", "No archiving conditions were found in the sources checked (see 'Sources checked' above)."])
        else:
            ws.append(["No.", "Condition", "Source", "Detail"])
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
                cell.fill = self._GREY
            for n, c in enumerate(res["conditions"], start=1):
                ws.append([n, c["condition"], c["source"], c["detail"]])
                for cell in ws[ws.max_row]:
                    cell.alignment = wrap
        ws.freeze_panes = None

    @staticmethod
    def _sheet_name(name: str, used: set) -> str:
        base = re.sub(r"[\\/*?:\[\]]", "_", name)[:31] or "Object"
        out, n = base, 2
        while out in used:
            out = f"{base[:28]}_{n}"
            n += 1
        used.add(out)
        return out

    @staticmethod
    def _save(wb) -> Path:
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        stem, suffix = Path(OUTPUT_NAME).stem, Path(OUTPUT_NAME).suffix
        candidates = [OUTPUT_DIR / OUTPUT_NAME] + [OUTPUT_DIR / f"{stem} ({n}){suffix}" for n in range(2, 10)]
        for path in candidates:
            try:
                wb.save(path)
                return path
            except PermissionError:   # open in Excel
                continue
        raise PermissionError("Could not write the analysis workbook (is it open in Excel?).")


job = ObjectAnalysisJob()
